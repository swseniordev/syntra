from io import BytesIO
from uuid import uuid4

from fastapi import HTTPException, Response
from openpyxl import load_workbook

from app.api.routers import mediciones
from app.core.identity import Principal


class _Connection:
    def __init__(self, row):
        self.row = row
        self.query = ""
        self.params = ()

    def execute(self, query, params):
        self.query, self.params = query, params
        return self

    def fetchone(self):
        return self.row


def _principal(permissions=("formatos:ver",)):
    return Principal(user={"id": 4}, role={"ver_todas": False},
                     permissions=frozenset(permissions), session_id=uuid4())


def _format_row():
    return {"id": "talkwalker_simple", "nombre": "Formato simple",
            "columnas": [
                {"campo_destino": "fecha_publicacion", "columnas_origen": ["published"], "transformacion": "fecha", "requerido": True},
                {"campo_destino": "texto", "columnas_origen": ["content", "title"], "transformacion": "coalesce", "requerido": True},
                {"campo_destino": "url", "columnas_origen": ["url"], "transformacion": "directo", "requerido": False},
            ]}


def test_format_template_issues_short_lived_user_and_session_bound_url(monkeypatch):
    signed = []
    monkeypatch.setattr(mediciones, "sign_format_template_download",
                        lambda user_id, session_id, format_id, digest: signed.append(
                            (user_id, session_id, format_id, digest)) or "signed-capability")
    response = Response()
    principal = _principal()
    result = mediciones.emitir_plantilla_formato("talkwalker_simple", response, None,
                                                 _Connection(_format_row()), principal)
    assert result["nombre"] == "plantilla_talkwalker_simple.xlsx"
    assert "token=signed-capability" in result["url"]
    assert result["expira_en_segundos"] == 900
    assert signed[0][:3] == (4, principal.session_id, "talkwalker_simple")
    assert len(signed[0][3]) == 64
    assert response.headers["Cache-Control"] == "private, no-store"


def test_format_template_download_is_real_xlsx_with_only_configured_headers(monkeypatch):
    row = _format_row()
    monkeypatch.setattr(mediciones, "verify_format_template_download",
                        lambda token, user_id, session_id, format_id, digest: token == "valid")
    response = Response()
    result = mediciones.emitir_plantilla_formato("talkwalker_simple", response, "valid",
                                                 _Connection(row), _principal())
    assert result.status_code == 200
    assert result.media_type == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    workbook = load_workbook(BytesIO(result.body), read_only=True)
    assert workbook.sheetnames == ["Plantilla"]
    assert next(workbook.active.iter_rows(values_only=True)) == ("published", "content", "title", "url")
    assert result.headers["Cache-Control"] == "private, no-store"


def test_format_template_rejects_missing_mapping_and_expired_or_wrong_capability(monkeypatch):
    with_missing_mapping = {"id": "meltwater", "nombre": "Meltwater", "columnas": []}
    with __import__("pytest").raises(HTTPException) as no_columns:
        mediciones.emitir_plantilla_formato("meltwater", Response(), None,
                                            _Connection(with_missing_mapping), _principal())
    assert no_columns.value.status_code == 409

    monkeypatch.setattr(mediciones, "verify_format_template_download", lambda *_args: False)
    with __import__("pytest").raises(HTTPException) as denied:
        mediciones.emitir_plantilla_formato("talkwalker_simple", Response(), "expired",
                                            _Connection(_format_row()), _principal())
    assert denied.value.status_code == 403

    with __import__("pytest").raises(HTTPException) as absent:
        mediciones.emitir_plantilla_formato("missing", Response(), None, _Connection(None), _principal())
    assert absent.value.status_code == 404
