"""Pruebas reales en clúster descartable; nunca escriben en syntra_dev.

Activación explícita: SYNTRA_RUN_DB_INTEGRATION=1 pytest -q tests/test_clients_integration.py
"""

import os
import json
import hashlib
from dataclasses import replace
from io import BytesIO
from pathlib import Path
import secrets
import shutil
import socket
import subprocess
import sys
import tempfile
import time
from datetime import datetime, timedelta, timezone
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier, Event

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import psycopg
from psycopg import sql
import pytest

from app.core.db import connection as production_connection
from app.core.security import password_hasher
from app.api.routers.auth import _ip_attempts, _ip_lock
from app.main import app

pytestmark = pytest.mark.skipif(os.getenv("SYNTRA_RUN_DB_INTEGRATION") != "1", reason="Requiere activar el clúster descartable de forma explícita")


@pytest.fixture(autouse=True)
def clear_test_ip_attempts():
    # Las pruebas comparten la IP artificial de TestClient; cada caso inicia una ventana nueva.
    with _ip_lock:
        _ip_attempts.clear()


@pytest.fixture(scope="module")
def environment():
    root = Path(__file__).resolve().parents[2]
    temp_root = root / "backend" / ".test-db"
    temp_root.mkdir(exist_ok=True)
    work = Path(tempfile.mkdtemp(prefix="syntra_integration_", dir=temp_root))
    data = work / "data"
    pg = Path("C:/Program Files/PostgreSQL/18/bin")
    with socket.socket() as sock:
        sock.bind(("127.0.0.1", 0))
        port = sock.getsockname()[1]

    def run(*args):
        background = Path(args[0]).name.lower() == "pg_ctl.exe"
        result = subprocess.run(
            args, stdin=subprocess.DEVNULL,
            stdout=subprocess.DEVNULL if background else subprocess.PIPE,
            stderr=subprocess.DEVNULL if background else subprocess.PIPE,
            text=True, encoding="utf-8", errors="replace", timeout=120, check=False,
        )
        if result.returncode:
            raise RuntimeError(f"Falló {Path(args[0]).name}: {(result.stderr or '')[-1600:]}")

    admin_password = secrets.token_urlsafe(24)
    started = False
    try:
        run(str(pg / "initdb.exe"), "-D", str(data), "-U", "postgres", "-A", "trust", "--encoding=UTF8", "--no-instructions")
        run(str(pg / "pg_ctl.exe"), "-D", str(data), "-l", str(work / "postgres.log"), "-o", f"-h 127.0.0.1 -p {port}", "-w", "start")
        started = True
        with psycopg.connect(host="127.0.0.1", port=port, dbname="postgres", user="postgres", autocommit=True) as superuser:
            superuser.execute("CREATE DATABASE syntra_trial")
        for script in (
            root / ".syntra-init-trial-20261020" / "ddl_trial.sql",
            root / ".syntra-init-trial-20261020" / "catalogos_tecnicos_trial.sql",
            root / "database" / "initialize_api_role.sql",
            root / "database" / "seed_rbac_oficial.sql",
            root / "database" / "entrega3_roles_privileges.sql",
            root / "database" / "entrega3_audit_read_privilege.sql",
            root / "database" / "entrega3_paises_privileges.sql",
            root / "database" / "entrega3_prompts_privileges.sql",
            root / "database" / "entrega3_regex_privileges.sql",
            root / "database" / "entrega3_cuentas_aprovisionamiento.sql",
            root / "database" / "entrega3_grupos_privileges.sql",
            root / "database" / "entrega3_usuarios_correo_privileges.sql",
            root / "database" / "entrega3_auditoria_exportacion_privileges.sql",
        ):
            run(str(pg / "psql.exe"), "-X", "-w", "-q", "-v", "ON_ERROR_STOP=1", "-h", "127.0.0.1", "-p", str(port), "-U", "postgres", "-d", "syntra_trial", "-f", str(script))
        # La migración se valida en este clúster temporal con un grantor no superusuario.
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres", autocommit=True) as superuser:
            superuser.execute("CREATE ROLE syntra_trial_deployer LOGIN CREATEROLE NOINHERIT NOSUPERUSER NOCREATEDB NOBYPASSRLS")
            superuser.execute("GRANT syntra_partition_owner TO syntra_trial_deployer")
            superuser.execute("GRANT CONNECT ON DATABASE syntra_trial TO syntra_trial_deployer WITH GRANT OPTION")
            superuser.execute("GRANT USAGE, CREATE ON SCHEMA ops,cat,cfg,med,prc,seg,aud TO syntra_trial_deployer WITH GRANT OPTION")
            superuser.execute("GRANT ALL PRIVILEGES ON ALL TABLES IN SCHEMA ops,cat,cfg,med,prc,seg,aud TO syntra_trial_deployer WITH GRANT OPTION")
            superuser.execute("GRANT EXECUTE ON ALL FUNCTIONS IN SCHEMA ops TO syntra_trial_deployer WITH GRANT OPTION")
            # El grantor de la migración debe poder retirar EXECUTE de PUBLIC
            # sobre la función oficial; en QA se ejecutará como su propietario.
            superuser.execute("ALTER FUNCTION ops.fn_preparar_medicion(integer,integer,integer[]) OWNER TO syntra_trial_deployer")
        run(str(pg / "psql.exe"), "-X", "-w", "-q", "-v", "ON_ERROR_STOP=1",
            "-v", "SLA_100K_S=180", "-v", "SLA_1M_S=1799", "-v", "MED_TIMEOUT_MS=1799000",
            "-h", "127.0.0.1", "-p", str(port), "-U", "syntra_trial_deployer", "-d", "syntra_trial",
            "-f", str(root / "database" / "entrega5_mediciones.sql"))
        run(str(pg / "psql.exe"), "-X", "-w", "-q", "-v", "ON_ERROR_STOP=1",
            "-h", "127.0.0.1", "-p", str(port), "-U", "postgres", "-d", "syntra_trial",
            "-f", str(root / "database" / "entrega5_mediciones_config_patch.sql"))
        for script in ("entrega5_catalogos_archivos.sql", "entrega5_catalogos_ingesta.sql",
                       "entrega5_talkwalker_formato_patch.sql",
                       "entrega5_formatos_api_grant.sql", "entrega5_formato_columnas_api_privileges.sql",
                       "entrega5_ingesta_api_grants.sql",
                       "entrega5_politica_ingesta.sql", "entrega5_ingesta_worker.sql",
                       "entrega5_ingesta_upsert_privileges.sql",
                       "entrega5_lotes_api_privileges.sql",
                       "entrega5_resumen_fuentes_api_privileges.sql",
                       "entrega5_notificaciones_ingesta_privileges.sql"):
            run(str(pg / "psql.exe"), "-X", "-w", "-q", "-v", "ON_ERROR_STOP=1",
                "-h", "127.0.0.1", "-p", str(port), "-U", "postgres", "-d", "syntra_trial",
                "-f", str(root / "database" / script))
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
            superuser.execute("ALTER ROLE syntra_api_login LOGIN")
            superuser.execute("ALTER ROLE syntra_mail_worker_login LOGIN")
            superuser.execute("ALTER ROLE syntra_report_worker_login LOGIN")
            superuser.execute("INSERT INTO seg.usuario(nombre,email,password_hash,rol_id) VALUES ('Prueba aislada','prueba@syntra.local',%s,1)", (password_hasher.hash(admin_password),))
            superuser.execute("INSERT INTO seg.usuario(nombre,email,password_hash,rol_id) VALUES ('Consultor aislado','consultor@syntra.local',%s,4)", (password_hasher.hash(admin_password),))

        def test_connection():
            with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="syntra_api_login", row_factory=psycopg.rows.dict_row) as conn:
                yield conn

        app.dependency_overrides[production_connection] = test_connection
        with TestClient(app, base_url="https://127.0.0.1:5173") as client:
            yield client, admin_password, port
    finally:
        app.dependency_overrides.clear()
        if started:
            run(str(pg / "pg_ctl.exe"), "-D", str(data), "-m", "immediate", "-w", "stop")
        if work.resolve().parent != temp_root.resolve() or not work.name.startswith("syntra_integration_"):
            raise RuntimeError("Ruta temporal inesperada; no se elimina")
        shutil.rmtree(work)


def _login(client, password, email="prueba@syntra.local"):
    result = client.post("/api/v1/auth/login", headers={"Origin": "https://127.0.0.1:5173"}, json={"email": email, "password": password})
    assert result.status_code == 200, result.text
    return {"Authorization": "Bearer " + result.json()["access_token"]}


def _deliver_trial_mail(port: int, process_id: int, *, fail: bool = False) -> str | None:
    """Ejecuta el worker real con un proveedor de prueba, sin SMTP ni secretos en disco."""
    code = """
import json
from app.config import WorkerSettings
from app.queue import MailQueue
class FakeProvider:
    sent = None
    def send_password(self, recipient, password, timeout, reason):
        if __import__('os').environ.get('SYNTRA_TEST_FAIL_MAIL') == '1':
            raise RuntimeError('Fallo de proveedor controlado')
        self.sent = (recipient, password)
provider = FakeProvider()
queue = MailQueue(WorkerSettings.from_environment(), provider)
task = queue.claim('trial-worker')
assert task is not None and task.process_id == int(__import__('os').environ['SYNTRA_TEST_PROCESS_ID'])
queue.process(task, 'trial-worker')
print(json.dumps({'password': provider.sent[1] if provider.sent else None}))
"""
    env = {**os.environ, "SYNTRA_WORKER_CONCURRENCY": "1", "SYNTRA_DB_HOST": "127.0.0.1",
           "SYNTRA_DB_PORT": str(port), "SYNTRA_DB_NAME": "syntra_trial",
           "SYNTRA_WORKER_DB_USER": "syntra_mail_worker_login", "SYNTRA_TEST_PROCESS_ID": str(process_id),
           "SYNTRA_TEST_FAIL_MAIL": "1" if fail else "0"}
    result = subprocess.run([sys.executable, "-c", code], cwd=Path(__file__).resolve().parents[2] / "worker",
                            env=env, text=True, capture_output=True, timeout=60, check=False)
    assert result.returncode == 0, result.stderr[-1200:]
    assert result.stdout.strip()
    return json.loads(result.stdout)["password"]


def test_usuarios_fallo_proveedor_reintento_y_auditoria(environment):
    client, password, port = environment
    admin = _login(client, password)
    created = client.post("/api/v1/usuarios", headers=admin,
                          json={"nombre": "Correo fallido", "email": "correo-fallido@syntra.local", "rol_id": 4})
    assert created.status_code == 201, created.text
    user_id = created.json()["id"]
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        process_id = superuser.execute("SELECT proceso_id FROM prc.proceso_tarea WHERE payload->>'usuario_id'=%s",
                                       (str(user_id),)).fetchone()[0]
    assert _deliver_trial_mail(port, process_id, fail=True) is None
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        task = superuser.execute("SELECT estado,payload,intentos FROM prc.proceso_tarea WHERE proceso_id=%s",
                                 (process_id,)).fetchone()
        assert task[0] == "en_cola" and task[2] == 1
        assert set(task[1]) == {"usuario_id", "version_esperada", "motivo"}
        assert superuser.execute("SELECT count(*) FROM aud.auditoria WHERE modulo='usuarios' "
                                 "AND accion='error_correo' AND entidad_id=%s", (str(user_id),)).fetchone()[0] == 1
        superuser.execute("UPDATE prc.proceso_tarea SET disponible_en=now() WHERE proceso_id=%s", (process_id,))
    delivered = _deliver_trial_mail(port, process_id)
    assert delivered
    assert _login(client, delivered, "correo-fallido@syntra.local")
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        assert superuser.execute("SELECT estado,intentos FROM prc.proceso_tarea WHERE proceso_id=%s",
                                 (process_id,)).fetchone() == ("completado", 2)


def test_usuarios_alta_cola_permisos_y_concurrencia(environment):
    client, password, port = environment
    admin = _login(client, password)
    consultor = _login(client, password, "consultor@syntra.local")
    assert client.post("/api/v1/usuarios", json={"nombre": "Sin acceso"}).status_code == 401
    assert client.post("/api/v1/usuarios", headers=consultor,
                       json={"nombre": "Sin acceso", "email": "sin@syntra.local", "rol_id": 4}).status_code == 403
    created = client.post("/api/v1/usuarios", headers=admin,
                          json={"nombre": "  Nuevo   Usuario ", "email": " NUEVO@SYNTRA.LOCAL ", "rol_id": 4})
    assert created.status_code == 201, created.text
    user_id = created.json()["id"]
    assert created.json()["nombre"] == "Nuevo Usuario"
    assert created.json()["email"] == "nuevo@syntra.local"
    assert created.json()["estado_id"] == 1
    assert "password_hash" not in created.json() and "version_token" not in created.json()
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        row = superuser.execute("SELECT password_hash,version_token FROM seg.usuario WHERE id=%s", (user_id,)).fetchone()
        assert row == (None, 1)
        task = superuser.execute("SELECT t.cola,t.estado,t.payload,p.tipo,p.id FROM prc.proceso_tarea t "
                                 "JOIN prc.proceso p ON p.id=t.proceso_id WHERE t.payload->>'usuario_id'=%s",
                                 (str(user_id),)).fetchone()
        assert task[0:2] == ("q.notificaciones", "en_cola")
        assert task[2] == {"usuario_id": user_id, "version_esperada": 1, "motivo": "alta"}
        assert task[3] == "correo_acceso"
    first_password = _deliver_trial_mail(port, task[4])
    first_login = _login(client, first_password, "nuevo@syntra.local")
    assert client.get("/api/v1/auth/yo", headers=first_login).status_code == 200
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        stored_hash = superuser.execute("SELECT password_hash FROM seg.usuario WHERE id=%s", (user_id,)).fetchone()[0]
        assert password_hasher.verify(stored_hash, first_password)
        assert stored_hash != first_password
    assert client.post("/api/v1/usuarios", headers=admin,
                       json={"nombre": "Duplicado", "email": "nuevo@syntra.local", "rol_id": 4}).status_code == 409
    detail = client.get(f"/api/v1/usuarios/{user_id}", headers=admin)
    assert detail.status_code == 200 and detail.headers["ETag"]
    assert client.patch(f"/api/v1/usuarios/{user_id}", headers=admin,
                        json={"telefono": "123"}).status_code == 428
    assert client.patch(f"/api/v1/usuarios/{user_id}", headers={**admin, "If-Match": '"stale"'},
                        json={"telefono": "123"}).status_code == 412
    changed = client.patch(f"/api/v1/usuarios/{user_id}", headers={**admin, "If-Match": detail.headers["ETag"]},
                           json={"telefono": "123"})
    assert changed.status_code == 200, changed.text
    assert changed.json()["telefono"] == "123"
    resend = client.post(f"/api/v1/usuarios/{user_id}/invitar", headers=admin)
    assert resend.status_code == 202, resend.text
    assert resend.json()["proceso_id"] > 0
    assert client.get("/api/v1/auth/yo", headers=first_login).status_code == 401
    second_password = _deliver_trial_mail(port, resend.json()["proceso_id"])
    assert second_password != first_password
    assert client.post("/api/v1/auth/login", headers={"Origin": "https://127.0.0.1:5173"},
                       json={"email": "nuevo@syntra.local", "password": first_password}).status_code == 401
    assert _login(client, second_password, "nuevo@syntra.local")
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        assert superuser.execute("SELECT version_token FROM seg.usuario WHERE id=%s", (user_id,)).fetchone()[0] == 4
        assert superuser.execute("SELECT count(*) FROM prc.proceso_tarea WHERE payload->>'usuario_id'=%s",
                                 (str(user_id),)).fetchone()[0] == 2
    inactive = client.post(f"/api/v1/usuarios/{user_id}/estado", headers=admin, json={"estado_id": 0})
    assert inactive.status_code == 200 and inactive.json()["estado_id"] == 0
    assert client.post(f"/api/v1/usuarios/{user_id}/invitar", headers=admin).status_code == 400
    assert client.delete(f"/api/v1/usuarios/{user_id}", headers=admin).status_code == 204
    restored = client.post(f"/api/v1/usuarios/{user_id}/restaurar", headers=admin)
    assert restored.status_code == 200 and restored.json()["estado_id"] == 0
    history = client.get(f"/api/v1/usuarios/{user_id}/historial", headers=admin)
    assert history.status_code == 200 and len(history.json()) >= 5


def test_clientes_persistencia_seguridad_y_concurrencia(environment):
    client, password, port = environment
    unauthenticated = client.get("/api/v1/clientes")
    assert unauthenticated.status_code == 401
    assert unauthenticated.headers["content-type"].startswith("application/problem+json")
    assert unauthenticated.json()["status"] == 401
    headers = _login(client, password)
    assert client.get("/api/v1/auth/yo", headers=headers).json()["permisos"]["clientes"] == ["crear", "editar", "eliminar", "ver"]
    empty = client.get("/api/v1/clientes", headers=headers)
    assert empty.status_code == 200 and empty.json()["items"] == []
    created = client.post("/api/v1/clientes", headers=headers, json={"nombre": "Cliente de ensayo", "sector": "Investigación", "estado_id": 1})
    assert created.status_code == 201, created.text
    identifier = created.json()["id"]
    assert created.json()["creado_por"] is not None
    detail = client.get(f"/api/v1/clientes/{identifier}", headers=headers)
    assert detail.status_code == 200 and detail.headers["ETag"]
    assert client.patch(f"/api/v1/clientes/{identifier}", headers=headers, json={"sector": "Nuevo"}).status_code == 428
    stale = client.patch(f"/api/v1/clientes/{identifier}", headers={**headers, "If-Match": '"stale"'}, json={"sector": "Nuevo"})
    assert stale.status_code == 412
    updated = client.patch(f"/api/v1/clientes/{identifier}", headers={**headers, "If-Match": detail.headers["ETag"]}, json={"sector": "Nuevo"})
    assert updated.status_code == 200, updated.text
    assert updated.headers["ETag"] != detail.headers["ETag"]
    assert client.get("/api/v1/clientes?q=Nuevo", headers=headers).json()["items"][0]["id"] == identifier
    assert client.post("/api/v1/clientes", headers=headers, json={"nombre": "Cliente de ensayo"}).status_code == 409
    assert client.post("/api/v1/clientes", headers=headers, json={"nombre": "Correo inválido", "contacto": "no-es-correo"}).status_code == 400
    assert client.patch(f"/api/v1/clientes/{identifier}", headers={**headers, "If-Match": updated.headers["ETag"]}, json={"contacto": "no-es-correo"}).status_code == 400
    assert client.get("/api/v1/clientes/999999", headers=headers).status_code == 404
    other = client.post("/api/v1/clientes", headers=headers, json={"nombre": "Segundo cliente", "sector": "Otro"})
    assert other.status_code == 201
    first_page = client.get("/api/v1/clientes?limite=1", headers=headers).json()
    assert len(first_page["items"]) == 1 and first_page["siguiente"]
    second_page = client.get("/api/v1/clientes?limite=1&cursor=" + first_page["siguiente"], headers=headers).json()
    assert len(second_page["items"]) == 1 and second_page["items"][0]["id"] == other.json()["id"]
    assert client.get("/api/v1/clientes?sector=Otro", headers=headers).json()["total_aprox"] == 1
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        assert superuser.execute("SELECT count(*) FROM cfg.cliente").fetchone()[0] == 2
        assert superuser.execute("SELECT count(*) FROM aud.auditoria WHERE modulo='clientes'").fetchone()[0] == 3
    assert client.patch(f"/api/v1/clientes/{identifier}", headers={**headers, "If-Match": updated.headers["ETag"]}, json={"estado_id": -1}).status_code == 400
    consultor = _login(client, password, "consultor@syntra.local")
    assert client.get("/api/v1/clientes", headers=consultor).status_code == 403
    assert client.post("/api/v1/clientes", headers=consultor, json={"nombre": "Prohibido"}).status_code == 403
    assert client.delete(f"/api/v1/clientes/{identifier}").status_code == 401
    assert client.delete(f"/api/v1/clientes/{identifier}", headers=consultor).status_code == 403

    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        account_id = superuser.execute(
            """INSERT INTO cfg.cuenta(cliente_id,nombre,prompt_clusters,prompt_temas,estado_id)
               VALUES (%s,'Cuenta de ensayo','clusters','temas',1) RETURNING id""", (identifier,),
        ).fetchone()[0]
    conflict = client.delete(f"/api/v1/clientes/{identifier}", headers=headers)
    assert conflict.status_code == 409
    assert conflict.json()["detail"] == "El cliente tiene cuentas asociadas activas. Para eliminarlo, primero debes desactivar las cuentas relacionadas"
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        assert superuser.execute("SELECT estado_id FROM cfg.cliente WHERE id=%s", (identifier,)).fetchone()[0] == 1
        assert superuser.execute("SELECT estado_id FROM cfg.cuenta WHERE id=%s", (account_id,)).fetchone()[0] == 1
        assert superuser.execute("SELECT count(*) FROM aud.auditoria WHERE modulo='clientes' AND entidad_id=%s AND accion='eliminar'", (str(identifier),)).fetchone()[0] == 0
        superuser.execute("UPDATE cfg.cuenta SET estado_id=0 WHERE id=%s", (account_id,))

    deleted = client.delete(f"/api/v1/clientes/{identifier}", headers=headers)
    assert deleted.status_code == 204, deleted.text
    assert client.delete(f"/api/v1/clientes/{identifier}", headers=headers).status_code == 404
    assert client.get(f"/api/v1/clientes/{identifier}", headers=headers).json()["estado_id"] == -1
    assert identifier not in [item["id"] for item in client.get("/api/v1/clientes?estado=vigentes", headers=headers).json()["items"]]
    assert identifier in [item["id"] for item in client.get("/api/v1/clientes?estado=eliminados", headers=headers).json()["items"]]
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        assert superuser.execute("SELECT estado_id FROM cfg.cliente WHERE id=%s", (identifier,)).fetchone()[0] == -1
        assert superuser.execute("SELECT estado_id FROM cfg.cuenta WHERE id=%s", (account_id,)).fetchone()[0] == 0
        event = superuser.execute("SELECT antes->>'estado_id', despues->>'estado_id' FROM aud.auditoria WHERE modulo='clientes' AND entidad_id=%s AND accion='eliminar'", (str(identifier),)).fetchone()
        assert event == ('1', '-1')


def test_cambio_estado_cliente(environment):
    client, password, port = environment
    headers = _login(client, password)
    created = client.post("/api/v1/clientes", headers=headers, json={"nombre": "Cliente de estados"})
    assert created.status_code == 201, created.text
    identifier = created.json()["id"]
    consultor = _login(client, password, "consultor@syntra.local")
    path = f"/api/v1/clientes/{identifier}/estado"
    assert client.post(path, json={"estado_id": 0}).status_code == 401
    assert client.post(path, headers=consultor, json={"estado_id": 0}).status_code == 403
    assert client.post(path, headers=headers, json={"estado_id": -1}).status_code == 400
    inactive = client.post(path, headers=headers, json={"estado_id": 0})
    assert inactive.status_code == 200 and inactive.json()["estado_id"] == 0
    assert identifier in [row["id"] for row in client.get("/api/v1/clientes?estado=inactivos", headers=headers).json()["items"]]
    assert identifier not in [row["id"] for row in client.get("/api/v1/clientes?estado=activos", headers=headers).json()["items"]]
    same = client.post(path, headers=headers, json={"estado_id": 0})
    assert same.status_code == 200 and same.json()["estado_id"] == 0
    active = client.post(path, headers=headers, json={"estado_id": 1})
    assert active.status_code == 200 and active.json()["estado_id"] == 1
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        assert superuser.execute("SELECT estado_id FROM cfg.cliente WHERE id=%s", (identifier,)).fetchone()[0] == 1
        events = superuser.execute(
            "SELECT antes->>'estado_id', despues->>'estado_id' FROM aud.auditoria WHERE modulo='clientes' AND entidad_id=%s AND accion='cambio_estado' ORDER BY id",
            (str(identifier),),
        ).fetchall()
        assert events == [('1', '0'), ('0', '1')]


def test_restauracion_conserva_cuenta_y_auditoria(environment):
    client, password, port = environment
    headers = _login(client, password)
    created = client.post("/api/v1/clientes", headers=headers, json={"nombre": "Cliente para restaurar"})
    assert created.status_code == 201, created.text
    identifier = created.json()["id"]
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        account_id = superuser.execute(
            """INSERT INTO cfg.cuenta(cliente_id,nombre,prompt_clusters,prompt_temas,estado_id)
               VALUES (%s,'Cuenta conservada','clusters','temas',0) RETURNING id""",
            (identifier,),
        ).fetchone()[0]
    assert client.delete(f"/api/v1/clientes/{identifier}", headers=headers).status_code == 204
    path = f"/api/v1/clientes/{identifier}/restaurar"
    assert client.post(path).status_code == 401
    consultor = _login(client, password, "consultor@syntra.local")
    assert client.post(path, headers=consultor).status_code == 403
    restored = client.post(path, headers=headers)
    assert restored.status_code == 200, restored.text
    assert restored.json()["id"] == identifier and restored.json()["estado_id"] == 0
    assert identifier in [r["id"] for r in client.get("/api/v1/clientes?estado=inactivos", headers=headers).json()["items"]]
    assert client.post(path, headers=headers).status_code == 400
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        assert superuser.execute("SELECT cliente_id,estado_id FROM cfg.cuenta WHERE id=%s", (account_id,)).fetchone() == (identifier, 0)
        events = superuser.execute(
            """SELECT accion,antes->>'estado_id',despues->>'estado_id'
               FROM aud.auditoria WHERE modulo='clientes' AND entidad_id=%s ORDER BY id""",
            (str(identifier),),
        ).fetchall()
        assert events == [('crear', None, '1'), ('eliminar', '1', '-1'), ('restaurar', '-1', '0')]

    conflict = client.post("/api/v1/clientes", headers=headers, json={"nombre": "Nombre para colisión"})
    assert conflict.status_code == 201
    conflict_id = conflict.json()["id"]
    assert client.delete(f"/api/v1/clientes/{conflict_id}", headers=headers).status_code == 204
    assert client.post("/api/v1/clientes", headers=headers, json={"nombre": "Nombre para colisión"}).status_code == 201
    assert client.post(f"/api/v1/clientes/{conflict_id}/restaurar", headers=headers).status_code == 400
    assert client.get(f"/api/v1/clientes/{conflict_id}", headers=headers).json()["estado_id"] == -1


def test_cuentas_asociadas_visibles_y_estado(environment):
    client, password, port = environment
    headers = _login(client, password)
    created = client.post("/api/v1/clientes", headers=headers, json={"nombre": "Cliente con cuentas"})
    assert created.status_code == 201
    identifier = created.json()["id"]
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        ids = [superuser.execute(
            """INSERT INTO cfg.cuenta(cliente_id,nombre,prompt_clusters,prompt_temas,estado_id)
               VALUES (%s,%s,'clusters','temas',%s) RETURNING id""",
            (identifier, name, state),
        ).fetchone()[0] for name, state in [('Cuenta activa aislada', 1), ('Cuenta inactiva aislada', 0), ('Cuenta eliminada aislada', -1)]]
    path = f"/api/v1/cuentas?cliente_id={identifier}"
    assert client.get(path).status_code == 401
    consultor = _login(client, password, "consultor@syntra.local")
    assert client.get(path, headers=consultor).json()["items"] == []
    result = client.get(path, headers=headers)
    assert result.status_code == 200, result.text
    assert [(r["id"], r["estado_id"]) for r in result.json()["items"]] == [(ids[0], 1), (ids[1], 0)]
    assert client.get(path + "&estado=activos", headers=headers).json()["items"][0]["id"] == ids[0]
    assert client.get(path + "&estado=inactivos", headers=headers).json()["items"][0]["id"] == ids[1]
    assert client.get(path + "&estado=eliminados", headers=headers).json()["items"][0]["id"] == ids[2]
    first = client.get(path + "&limite=1", headers=headers).json()
    second = client.get(path + "&limite=1&cursor=" + first["siguiente"], headers=headers).json()
    assert first["items"][0]["id"] == ids[0] and second["items"][0]["id"] == ids[1]
    assert client.get(path + "&limite=1&cursor=erroneo", headers=headers).status_code == 400

    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        group_id = superuser.execute("INSERT INTO seg.grupo(nombre) VALUES ('Grupo aislado') RETURNING id").fetchone()[0]
        user_id = superuser.execute("SELECT id FROM seg.usuario WHERE email='consultor@syntra.local'").fetchone()[0]
        superuser.execute("INSERT INTO seg.grupo_usuario(grupo_id,usuario_id) VALUES (%s,%s)", (group_id, user_id))
        superuser.execute("INSERT INTO seg.grupo_cuenta(grupo_id,cuenta_id) VALUES (%s,%s)", (group_id, ids[1]))
    visible = client.get(path, headers=consultor)
    assert visible.status_code == 200 and [r["id"] for r in visible.json()["items"]] == [ids[1]]


def test_eliminacion_espera_activacion_concurrente(environment):
    client, password, port = environment
    headers = _login(client, password)
    created = client.post("/api/v1/clientes", headers=headers, json={"nombre": "Cliente concurrente"})
    assert created.status_code == 201, created.text
    identifier = created.json()["id"]
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        account_id = superuser.execute(
            """INSERT INTO cfg.cuenta(cliente_id,nombre,prompt_clusters,prompt_temas,estado_id)
               VALUES (%s,'Cuenta concurrente','clusters','temas',0) RETURNING id""", (identifier,),
        ).fetchone()[0]
    entered = Event()

    def delete_from_other_thread():
        entered.set()
        with TestClient(app, base_url="https://127.0.0.1:5173") as second_client:
            return second_client.delete(f"/api/v1/clientes/{identifier}", headers=headers)

    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as activation:
        activation.execute("UPDATE cfg.cuenta SET estado_id=1 WHERE id=%s", (account_id,))
        with ThreadPoolExecutor(max_workers=1) as pool:
            result = pool.submit(delete_from_other_thread)
            assert entered.wait(timeout=5)
            deadline = time.monotonic() + 5
            blocked = False
            with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres", autocommit=True) as observer:
                while time.monotonic() < deadline:
                    blocked = observer.execute(
                        """SELECT EXISTS (
                             SELECT 1 FROM pg_stat_activity
                             WHERE datname='syntra_trial' AND usename='syntra_api_login'
                               AND wait_event_type='Lock'
                               AND query LIKE 'SELECT estado_id FROM cfg.cuenta%')""",
                    ).fetchone()[0]
                    if blocked:
                        break
                    time.sleep(0.05)
            assert blocked, "La eliminación no esperó el bloqueo de la cuenta"
            activation.commit()
            conflict = result.result(timeout=10)
    assert conflict.status_code == 409, conflict.text
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        assert superuser.execute("SELECT estado_id FROM cfg.cliente WHERE id=%s", (identifier,)).fetchone()[0] == 1
        assert superuser.execute("SELECT estado_id FROM cfg.cuenta WHERE id=%s", (account_id,)).fetchone()[0] == 1


def test_refresh_rotacion_reutilizacion_y_logout(environment):
    client, password, port = environment
    headers = _login(client, password)
    old_cookie = client.cookies.get("refresh")
    renewed = client.post("/api/v1/auth/refresh", headers={"Origin": "https://127.0.0.1:5173"})
    assert renewed.status_code == 200, renewed.text
    assert client.get("/api/v1/auth/yo", headers=headers).status_code == 401
    with TestClient(app, base_url="https://127.0.0.1:5173") as attacker:
        attacker.cookies.set("refresh", old_cookie, path="/api/v1/auth")
        replay = attacker.post("/api/v1/auth/refresh", headers={"Origin": "https://127.0.0.1:5173"})
        assert replay.status_code == 401
    assert client.post("/api/v1/auth/refresh", headers={"Origin": "https://127.0.0.1:5173"}).status_code == 401
    fresh = _login(client, password)
    assert client.post("/api/v1/auth/logout", headers={**fresh, "Origin": "https://127.0.0.1:5173"}).status_code == 204
    assert client.get("/api/v1/auth/yo", headers=fresh).status_code == 401
    again = _login(client, password)
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        superuser.execute("UPDATE seg.usuario SET version_token=version_token+1 WHERE email='prueba@syntra.local'")
    assert client.get("/api/v1/auth/yo", headers=again).status_code == 401
    assert client.post("/api/v1/auth/refresh", headers={"Origin": "https://127.0.0.1:5173"}).status_code == 401


def test_refresh_concurrente_consume_una_vez_y_revoca_familia(environment):
    client, password, port = environment
    _login(client, password)
    original = client.cookies.get("refresh")
    assert original
    sid = original.split(".", 1)[0]
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as observer:
        family = observer.execute("SELECT familia_id FROM seg.sesion WHERE id=%s", (sid,)).fetchone()[0]

    start = Barrier(3)

    def consume_original():
        with TestClient(app, base_url="https://127.0.0.1:5173") as contender:
            contender.cookies.set("refresh", original, path="/api/v1/auth")
            start.wait(timeout=10)
            result = contender.post("/api/v1/auth/refresh", headers={"Origin": "https://127.0.0.1:5173"})
            return result.status_code, result.json().get("access_token") if result.status_code == 200 else None

    with ThreadPoolExecutor(max_workers=2) as pool:
        first = pool.submit(consume_original)
        second = pool.submit(consume_original)
        start.wait(timeout=10)
        outcomes = [first.result(timeout=20), second.result(timeout=20)]
    assert sorted(status for status, _ in outcomes) == [200, 401]
    winner = next(token for status, token in outcomes if status == 200)
    assert client.get("/api/v1/auth/yo", headers={"Authorization": f"Bearer {winner}"}).status_code == 401
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as observer:
        active = observer.execute("SELECT count(*) FROM seg.sesion WHERE familia_id=%s AND revocado_en IS NULL", (family,)).fetchone()[0]
    assert active == 0

def test_normalizacion_sectores_y_autoria_real(environment):
    client, password, port = environment
    admin = _login(client, password)
    assert client.get("/api/v1/clientes/sectores").status_code == 401
    assert client.get("/api/v1/clientes/sectores", headers=_login(client, password, "consultor@syntra.local")).status_code == 403

    created = client.post("/api/v1/clientes", headers=admin, json={
        "nombre": "  Cliente\t normalizado\n", "rfc": "  RFC123\n", "sector": "  operación  ",
        "contacto": "  prueba@example.test\n",
    })
    assert created.status_code == 201, created.text
    record = created.json()
    assert (record["nombre"], record["rfc"], record["sector"], record["contacto"]) == (
        "Cliente normalizado", "RFC123", "OPERACIÓN", "prueba@example.test",
    )
    detail = client.get(f"/api/v1/clientes/{record['id']}", headers=admin)
    assert detail.status_code == 200, detail.text
    assert detail.json()["creado_por_nombre"] == "Prueba aislada"
    assert detail.json()["actualizado_por_nombre"] == "Prueba aislada"
    original_created_at = detail.json()["creado_en"]

    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        superuser.execute("INSERT INTO seg.usuario(nombre,email,password_hash,rol_id) VALUES ('Editor aislado','editor@syntra.local',%s,1)", (password_hasher.hash(password),))
        superuser.execute("INSERT INTO cfg.cliente(nombre,sector,estado_id) VALUES ('Sector legado uno',' Operacion ',1), ('Sector legado dos','OPERACION',0), ('Sector legado tres',' operacion  ',-1), ('Sector acentuado','  operación  ',1), ('Sin sector','  ',1), ('  Nombre legado  ',NULL,1)")
    editor = _login(client, password, "editor@syntra.local")
    changed = client.patch(f"/api/v1/clientes/{record['id']}", headers={**editor, "If-Match": detail.headers["ETag"]}, json={
        "nombre": "  Cliente\n  normalizado\t", "rfc": "\n  RFC  456\tB  ", "sector": "  operacion \n", "contacto": "  nuevo@example.test  ",
    })
    assert changed.status_code == 200, changed.text
    assert (changed.json()["nombre"], changed.json()["rfc"], changed.json()["sector"], changed.json()["contacto"]) == (
        "Cliente normalizado", "RFC  456\tB", "OPERACION", "nuevo@example.test",
    )
    assert changed.json()["creado_en"] == original_created_at
    refreshed = client.get(f"/api/v1/clientes/{record['id']}", headers=admin).json()
    assert refreshed["creado_por_nombre"] == "Prueba aislada"
    assert refreshed["actualizado_por_nombre"] == "Editor aislado"
    assert refreshed["creado_en"] == original_created_at
    assert refreshed["actualizado_en"] > original_created_at

    sectors = client.get("/api/v1/clientes/sectores", headers=admin)
    assert sectors.status_code == 200, sectors.text
    assert sectors.json()["items"].count("OPERACION") == 1
    assert "OPERACIÓN" in sectors.json()["items"]
    assert "" not in sectors.json()["items"]
    filtered = client.get("/api/v1/clientes?sector= operacion &limite=1", headers=admin).json()
    assert filtered["total_aprox"] == 3
    assert filtered["siguiente"]
    assert client.get("/api/v1/clientes?sector= operacion &limite=1&cursor=" + filtered["siguiente"], headers=admin).status_code == 200
    assert client.get("/api/v1/clientes?sector= OPERACIÓN ", headers=admin).json()["total_aprox"] == 1
    assert client.post("/api/v1/clientes", headers=admin, json={"nombre": "Nombre legado"}).status_code == 409
    assert client.post("/api/v1/clientes", headers=admin, json={"nombre": "  Cliente normalizado  "}).status_code == 409
    assert client.post("/api/v1/clientes", headers=admin, json={"nombre": "  \n  "}).status_code == 400
    assert client.patch(f"/api/v1/clientes/{record['id']}", headers={**editor, "If-Match": client.get(f"/api/v1/clientes/{record['id']}", headers=editor).headers["ETag"]}, json={"nombre": "Nombre legado"}).status_code == 409
    current = client.get(f"/api/v1/clientes/{record['id']}", headers=editor)
    cleared = client.patch(f"/api/v1/clientes/{record['id']}", headers={**editor, "If-Match": current.headers["ETag"]}, json={"rfc": "\n ", "sector": "\t", "contacto": "  "})
    assert cleared.status_code == 200, cleared.text
    assert (cleared.json()["rfc"], cleared.json()["sector"], cleared.json()["contacto"]) == (None, None, None)
    internal = client.post("/api/v1/clientes", headers=admin, json={
        "nombre": " Cliente  con\t espacios\n", "sector": "  acción\t  social\n",
        "rfc": " \nRFC  12\tB\n ", "contacto": "  correo@example.test \n",
    })
    assert internal.status_code == 201, internal.text
    assert (internal.json()["nombre"], internal.json()["sector"], internal.json()["rfc"], internal.json()["contacto"]) == (
        "Cliente con espacios", "ACCIÓN SOCIAL", "RFC  12\tB", "correo@example.test",
    )
    assert client.post("/api/v1/clientes", headers=admin, json={"nombre": "Cliente   con\n espacios"}).status_code == 409
    assert client.post("/api/v1/clientes", headers=admin, json={"nombre": "Cliente con espacios", "contacto": "correo @example.test"}).status_code == 400
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        superuser.execute("INSERT INTO cfg.cliente(nombre,estado_id) VALUES ('Nombre  legado con  espacios',1)")
    assert client.post("/api/v1/clientes", headers=admin, json={"nombre": "Nombre legado\t con espacios"}).status_code == 409
    fresh = client.get(f"/api/v1/clientes/{internal.json()['id']}", headers=editor)
    assert client.patch(f"/api/v1/clientes/{internal.json()['id']}", headers={**editor, "If-Match": fresh.headers["ETag"]}, json={
        "nombre": "Nombre\t legado  con\n espacios",
    }).status_code == 409
    normalized_patch = client.patch(
        f"/api/v1/clientes/{internal.json()['id']}",
        headers={**editor, "If-Match": fresh.headers["ETag"]},
        json={"nombre": " Cliente\t con\n espacios ", "sector": "  análisis\t  político\n"},
    )
    assert normalized_patch.status_code == 200, normalized_patch.text
    assert (normalized_patch.json()["nombre"], normalized_patch.json()["sector"]) == ("Cliente con espacios", "ANÁLISIS POLÍTICO")
    persisted = client.get(f"/api/v1/clientes/{internal.json()['id']}", headers=admin)
    assert (persisted.json()["nombre"], persisted.json()["sector"]) == ("Cliente con espacios", "ANÁLISIS POLÍTICO")
    optional = client.post("/api/v1/clientes", headers=admin, json={
        "nombre": "Opcionales vacíos", "rfc": "\n \t", "sector": "  \n", "contacto": " \t ",
    })
    assert optional.status_code == 201, optional.text
    assert (optional.json()["rfc"], optional.json()["sector"], optional.json()["contacto"]) == (None, None, None)


def test_fallo_de_auditoria_revierte_modificacion_cliente(environment):
    client, password, port = environment
    headers = _login(client, password)
    created = client.post("/api/v1/clientes", headers=headers, json={"nombre": "Cliente rollback auditoría", "sector": "ORIGINAL"})
    assert created.status_code == 201, created.text
    identifier = created.json()["id"]
    initial_etag = created.headers["ETag"]
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        before_count = superuser.execute(
            "SELECT count(*) FROM aud.auditoria WHERE modulo='clientes' AND entidad_id=%s", (str(identifier),),
        ).fetchone()[0]
        superuser.execute("REVOKE INSERT ON aud.auditoria FROM syntra_api_login")
    try:
        with TestClient(app, base_url="https://127.0.0.1:5173", raise_server_exceptions=False) as failing_client:
            result = failing_client.patch(
                f"/api/v1/clientes/{identifier}",
                headers={**headers, "If-Match": initial_etag},
                json={"sector": "CAMBIADO"},
            )
        assert result.status_code == 500
    finally:
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
            superuser.execute("GRANT INSERT ON aud.auditoria TO syntra_api_login")
    current = client.get(f"/api/v1/clientes/{identifier}", headers=headers)
    assert current.status_code == 200
    assert current.json()["sector"] == "ORIGINAL"
    assert current.headers["ETag"] == initial_etag
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        assert superuser.execute(
            "SELECT count(*) FROM aud.auditoria WHERE modulo='clientes' AND entidad_id=%s", (str(identifier),),
        ).fetchone()[0] == before_count


def test_roles_persistencia_permisos_y_auditoria(environment):
    client, password, port = environment
    admin_headers = _login(client, password)
    consultor_headers = _login(client, password, "consultor@syntra.local")
    assert client.get("/api/v1/roles", headers=consultor_headers).status_code == 403
    assert client.post("/api/v1/roles", headers=consultor_headers, json={
        "nombre": "No autorizado", "permisos": {},
    }).status_code == 403

    page = client.get("/api/v1/roles", headers=admin_headers)
    assert page.status_code == 200 and page.json()["total_aprox"] == 5
    created = client.post("/api/v1/roles", headers=admin_headers, json={
        "nombre": "  Equipo   controlado  ", "descripcion": "Rol de prueba aislada",
        "ver_todas": False, "permisos": {"clientes": ["ver"]}, "estado_id": 1,
    })
    assert created.status_code == 201, created.text
    role_id = created.json()["id"]
    assert created.json()["nombre"] == "Equipo controlado"
    assert created.json()["permisos"] == {"clientes": ["ver"]}

    detail = client.get(f"/api/v1/roles/{role_id}", headers=admin_headers)
    assert detail.status_code == 200 and detail.headers["ETag"]
    stale = client.patch(f"/api/v1/roles/{role_id}", headers={**admin_headers, "If-Match": '"stale"'},
                         json={"nombre": "Otro nombre"})
    assert stale.status_code == 412
    edited = client.patch(f"/api/v1/roles/{role_id}",
                          headers={**admin_headers, "If-Match": detail.headers["ETag"]},
                          json={"nombre": "Equipo revisado"})
    assert edited.status_code == 200, edited.text
    assert client.get(f"/api/v1/roles/{role_id}", headers=admin_headers).json()["nombre"] == "Equipo revisado"
    history = client.get(f"/api/v1/roles/{role_id}/historial", headers=admin_headers)
    assert history.status_code == 200, history.text
    assert {event["accion"] for event in history.json()} == {"crear", "editar"}

    own_role = client.get("/api/v1/roles/1", headers=admin_headers)
    self_lockout = client.patch("/api/v1/roles/1", headers={**admin_headers, "If-Match": own_role.headers["ETag"]},
                                json={"permisos": {}})
    assert self_lockout.status_code == 409
    assert client.get("/api/v1/roles/1", headers=admin_headers).json()["permisos"]["roles"]

    deleted = client.delete(f"/api/v1/roles/{role_id}", headers=admin_headers)
    assert deleted.status_code == 204
    assert client.get(f"/api/v1/roles/{role_id}", headers=admin_headers).json()["estado_id"] == -1
    restored = client.post(f"/api/v1/roles/{role_id}/restaurar", headers=admin_headers)
    assert restored.status_code == 200 and restored.json()["estado_id"] == 0
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        assert superuser.execute(
            "SELECT count(*) FROM aud.auditoria WHERE modulo='roles' AND entidad_id=%s", (str(role_id),),
        ).fetchone()[0] == 4


def test_paises_persistencia_permisos_y_auditoria(environment):
    client, password, port = environment
    admin = _login(client, password)
    consultor = _login(client, password, "consultor@syntra.local")
    assert client.get("/api/v1/paises").status_code == 401
    assert client.get("/api/v1/paises", headers=consultor).status_code == 403
    assert client.post("/api/v1/paises", headers=consultor, json={"nombre": "No autorizado", "codigo": "XX"}).status_code == 403
    assert client.get("/api/v1/paises", headers=admin).json()["items"] == []

    created = client.post("/api/v1/paises", headers=admin,
                          json={"nombre": "  Costa   Rica ", "codigo": "CR", "idioma": "es-419", "estado_id": 1})
    assert created.status_code == 201, created.text
    country_id = created.json()["id"]
    assert created.json()["nombre"] == "Costa Rica"
    assert created.headers["ETag"] and created.headers["Location"].endswith(f"/{country_id}")
    assert client.post("/api/v1/paises", headers=admin, json={"nombre": "Duplicado", "codigo": "CR"}).status_code == 409
    detail = client.get(f"/api/v1/paises/{country_id}", headers=admin)
    assert detail.status_code == 200 and detail.headers["ETag"] == created.headers["ETag"]
    assert client.patch(f"/api/v1/paises/{country_id}", headers=admin, json={"idioma": "es"}).status_code == 428
    assert client.patch(f"/api/v1/paises/{country_id}", headers={**admin, "If-Match": '"stale"'},
                        json={"idioma": "es"}).status_code == 412
    edited = client.patch(f"/api/v1/paises/{country_id}", headers={**admin, "If-Match": detail.headers["ETag"]},
                          json={"idioma": "es"})
    assert edited.status_code == 200, edited.text
    assert edited.json()["idioma"] == "es" and edited.headers["ETag"] != detail.headers["ETag"]
    assert client.get("/api/v1/paises?q=Costa", headers=admin).json()["items"][0]["id"] == country_id
    assert client.post(f"/api/v1/paises/{country_id}/estado", headers=admin,
                       json={"estado_id": 0}).json()["estado_id"] == 0
    assert client.delete(f"/api/v1/paises/{country_id}", headers=admin).status_code == 204
    assert client.get(f"/api/v1/paises/{country_id}", headers=admin).json()["estado_id"] == -1
    assert client.post(f"/api/v1/paises/{country_id}/restaurar", headers=admin).json()["estado_id"] == 0
    history = client.get(f"/api/v1/paises/{country_id}/historial", headers=admin)
    assert history.status_code == 200, history.text
    assert {event["accion"] for event in history.json()} == {"crear", "editar", "cambio_estado", "eliminar", "restaurar"}
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        assert superuser.execute("SELECT count(*) FROM cfg.pais WHERE id=%s", (country_id,)).fetchone()[0] == 1
        assert superuser.execute("SELECT count(*) FROM aud.auditoria WHERE modulo='paises' AND entidad_id=%s",
                                 (str(country_id),)).fetchone()[0] == 5


def test_prompts_versiones_predeterminado_y_permisos(environment):
    client, password, port = environment
    admin = _login(client, password)
    consultor = _login(client, password, "consultor@syntra.local")
    assert client.get("/api/v1/prompts").status_code == 401
    assert client.get("/api/v1/prompts", headers=consultor).status_code == 403
    first = client.post("/api/v1/prompts", headers=admin, json={
        "tipo": "cluster", "red": None, "nombre": "Primer prompt",
        "contenido": "Agrupa los registros en categorías políticas.", "por_defecto": True,
    })
    assert first.status_code == 201, first.text
    first_id = first.json()["id"]
    second = client.post("/api/v1/prompts", headers=admin, json={
        "tipo": "cluster", "red": None, "nombre": "Segundo prompt",
        "contenido": "Agrupa los registros por asunto público.", "por_defecto": True,
    })
    assert second.status_code == 201, second.text
    second_id = second.json()["id"]
    assert client.get(f"/api/v1/prompts/{first_id}", headers=admin).json()["por_defecto"] is False
    assert client.get(f"/api/v1/prompts/{second_id}", headers=admin).json()["por_defecto"] is True
    assert client.patch(f"/api/v1/prompts/{second_id}", headers=admin, json={"nombre": "Sin ETag"}).status_code == 428
    assert client.patch(f"/api/v1/prompts/{second_id}", headers={**admin, "If-Match": '"stale"'},
                        json={"nombre": "Con ETag obsoleto"}).status_code == 412
    edited = client.patch(f"/api/v1/prompts/{second_id}",
                          headers={**admin, "If-Match": second.headers["ETag"]},
                          json={"contenido": "Agrupa las publicaciones por asunto político actualizado."})
    assert edited.status_code == 200, edited.text
    assert edited.json()["version_actual"] == 2
    versions = client.get(f"/api/v1/prompts/{second_id}/versiones", headers=admin)
    assert versions.status_code == 200 and [row["version"] for row in versions.json()] == [2, 1]
    assert versions.json()[1]["contenido"] == second.json()["contenido"]
    assert client.get("/api/v1/prompts?tipo=cluster&red=_", headers=admin).json()["total_aprox"] == 2
    inactive = client.post(f"/api/v1/prompts/{second_id}/estado", headers=admin, json={"estado_id": 0})
    assert inactive.status_code == 200 and inactive.json()["por_defecto"] is False
    assert client.delete(f"/api/v1/prompts/{second_id}", headers=admin).status_code == 204
    assert client.post(f"/api/v1/prompts/{second_id}/restaurar", headers=admin).json()["estado_id"] == 0
    history = client.get(f"/api/v1/prompts/{second_id}/historial", headers=admin)
    assert history.status_code == 200 and len(history.json()) == 5
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        assert superuser.execute("SELECT count(*) FROM cfg.prompt_version WHERE prompt_id=%s", (second_id,)).fetchone()[0] == 2
        assert superuser.execute("SELECT count(*) FROM cfg.prompt WHERE tipo='cluster' AND por_defecto AND estado_id=1").fetchone()[0] == 0


def test_regex_re2_versiones_permisos_y_auditoria(environment):
    client, password, port = environment
    admin = _login(client, password)
    consultor = _login(client, password, "consultor@syntra.local")
    assert client.get("/api/v1/regex").status_code == 401
    assert client.get("/api/v1/regex", headers=consultor).status_code == 403
    assert client.post("/api/v1/regex/probar", headers=consultor,
                       json={"condicion": "agua", "texto": "agua"}).status_code == 403
    tested = client.post("/api/v1/regex/probar", headers=admin,
                         json={"condicion": "agua", "flags": "i", "texto": "Agua y agua"})
    assert tested.status_code == 200 and len(tested.json()["coincidencias"]) == 2
    invalid = client.post("/api/v1/regex/probar", headers=admin,
                          json={"condicion": "(?=agua)agua", "texto": "agua"})
    assert invalid.status_code == 200 and invalid.json()["valida"] is False
    assert client.post("/api/v1/regex", headers=admin, json={
        "nombre": "No cumple", "condicion": "agua", "texto_prueba": "sequía",
    }).status_code == 400
    assert client.post("/api/v1/regex", headers=admin, json={
        "nombre": "RE2 inválida", "condicion": "(?=agua)agua", "texto_prueba": "agua",
    }).status_code == 400

    created = client.post("/api/v1/regex", headers=admin, json={
        "nombre": "  Crisis   de agua ", "condicion": "agua", "flags": "i",
        "texto_prueba": "Agua en la ciudad", "descripcion": "Detección de agua",
    })
    assert created.status_code == 201, created.text
    regex_id = created.json()["id"]
    assert created.json()["nombre"] == "Crisis de agua"
    assert created.json()["version"] == 1
    assert created.headers["ETag"]
    assert client.post("/api/v1/regex", headers=admin, json={
        "nombre": "Crisis de agua", "condicion": "agua", "texto_prueba": "agua",
    }).status_code == 409
    assert client.get(f"/api/v1/regex/{regex_id}", headers=admin).status_code == 200
    assert client.get("/api/v1/regex?q=Crisis", headers=admin).json()["total_aprox"] == 1
    assert client.patch(f"/api/v1/regex/{regex_id}", headers=admin,
                        json={"descripcion": "Actualizada"}).status_code == 428
    assert client.patch(f"/api/v1/regex/{regex_id}", headers={**admin, "If-Match": '"stale"'},
                        json={"descripcion": "Actualizada"}).status_code == 412
    metadata = client.patch(f"/api/v1/regex/{regex_id}",
                            headers={**admin, "If-Match": created.headers["ETag"]},
                            json={"descripcion": "Actualizada"})
    assert metadata.status_code == 200 and metadata.json()["version"] == 1
    changed = client.patch(f"/api/v1/regex/{regex_id}",
                           headers={**admin, "If-Match": metadata.headers["ETag"]},
                           json={"condicion": "agua|sequía", "texto_prueba": "Hay sequía"})
    assert changed.status_code == 200, changed.text
    assert changed.json()["version"] == 2
    assert client.post(f"/api/v1/regex/{regex_id}/estado", headers=admin,
                       json={"estado_id": 0}).json()["estado_id"] == 0
    assert client.delete(f"/api/v1/regex/{regex_id}", headers=admin).status_code == 204
    assert client.post(f"/api/v1/regex/{regex_id}/restaurar", headers=admin).json()["estado_id"] == 0
    history = client.get(f"/api/v1/regex/{regex_id}/historial", headers=admin)
    assert history.status_code == 200 and len(history.json()) == 6
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        versions = superuser.execute(
            "SELECT version,condicion FROM cfg.regex_version WHERE regex_id=%s ORDER BY version", (regex_id,),
        ).fetchall()
        assert versions == [(1, "agua"), (2, "agua|sequía")]
        assert superuser.execute("SELECT count(*) FROM aud.auditoria WHERE modulo='regex' AND entidad_id=%s",
                                 (str(regex_id),)).fetchone()[0] == 6


def test_cuentas_aprovisionamiento_privilegios_y_rollback(environment):
    _client, _password, port = environment
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        customer_id = superuser.execute(
            "INSERT INTO cfg.cliente(nombre) VALUES ('Cliente para cuentas aisladas') RETURNING id"
        ).fetchone()[0]
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial",
                         user="syntra_api_login", row_factory=psycopg.rows.dict_row) as api_conn:
        assert api_conn.execute("SELECT has_schema_privilege(current_user,'dat','CREATE') AS allowed").fetchone()["allowed"] is False
        assert api_conn.execute("SELECT has_function_privilege(current_user,'ops.fn_preparar_cuenta(integer)','EXECUTE') AS allowed").fetchone()["allowed"] is False
        assert api_conn.execute("SELECT has_function_privilege(current_user,'ops.fn_crear_y_adjuntar(text,text,integer,text,text)','EXECUTE') AS allowed").fetchone()["allowed"] is False

        rolled_back = api_conn.execute(
            "INSERT INTO cfg.cuenta(cliente_id,nombre,prompt_clusters,prompt_temas,creado_por,actualizado_por) "
            "VALUES (%s,'Cuenta revertida','Agrupa por tema','Resume temas',1,1) RETURNING id", (customer_id,),
        ).fetchone()["id"]
        api_conn.execute("SELECT ops.fn_aprovisionar_cuenta(%s,1)", (rolled_back,))
        api_conn.execute(
            "INSERT INTO aud.auditoria(usuario_id,modulo,accion,entidad_tipo,entidad_id) "
            "VALUES (1,'cuentas','crear','Cuenta',%s)", (str(rolled_back),),
        )
        api_conn.rollback()

        identifier = api_conn.execute(
            "INSERT INTO cfg.cuenta(cliente_id,nombre,prompt_clusters,prompt_temas,creado_por,actualizado_por) "
            "VALUES (%s,'Cuenta completa','Agrupa por tema','Resume temas',1,1) RETURNING id", (customer_id,),
        ).fetchone()["id"]
        api_conn.execute("SELECT ops.fn_aprovisionar_cuenta(%s,1)", (identifier,))
        api_conn.execute(
            "INSERT INTO aud.auditoria(usuario_id,modulo,accion,entidad_tipo,entidad_id) "
            "VALUES (1,'cuentas','crear','Cuenta',%s)", (str(identifier),),
        )
        api_conn.commit()
        with pytest.raises(psycopg.Error):
            api_conn.execute("SELECT ops.fn_aprovisionar_cuenta(%s,1)", (identifier,))
        api_conn.rollback()

    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        assert superuser.execute("SELECT count(*) FROM cfg.cuenta WHERE id=%s", (rolled_back,)).fetchone()[0] == 0
        assert superuser.execute("SELECT count(*) FROM aud.auditoria WHERE modulo='cuentas' AND entidad_id=%s",
                                 (str(rolled_back),)).fetchone()[0] == 0
        assert superuser.execute(
            "SELECT count(*) FROM pg_inherits i JOIN pg_class child ON child.oid=i.inhrelid "
            "JOIN pg_class parent ON parent.oid=i.inhparent JOIN pg_namespace n ON n.oid=parent.relnamespace "
            "WHERE n.nspname='dat' AND child.relname LIKE %s", (f"%_c{rolled_back}",),
        ).fetchone()[0] == 0
        assert superuser.execute("SELECT count(*) FROM cfg.cuenta WHERE id=%s", (identifier,)).fetchone()[0] == 1
        assert superuser.execute("SELECT count(*) FROM aud.auditoria WHERE modulo='cuentas' AND entidad_id=%s",
                                 (str(identifier),)).fetchone()[0] == 1
        assert superuser.execute(
            "SELECT count(*) FROM pg_inherits i JOIN pg_class child ON child.oid=i.inhrelid "
            "JOIN pg_class parent ON parent.oid=i.inhparent JOIN pg_namespace n ON n.oid=parent.relnamespace "
            "WHERE n.nspname='dat' AND child.relname LIKE %s", (f"%_c{identifier}",),
        ).fetchone()[0] == 12


def test_cuentas_fallo_de_particion_revierte_alta_y_evento(environment):
    _client, _password, port = environment
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        customer_id = superuser.execute(
            "INSERT INTO cfg.cliente(nombre) VALUES ('Cliente para fallo aislado') RETURNING id"
        ).fetchone()[0]
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial",
                         user="syntra_api_login", row_factory=psycopg.rows.dict_row) as api_conn:
        identifier = api_conn.execute(
            "INSERT INTO cfg.cuenta(cliente_id,nombre,prompt_clusters,prompt_temas,creado_por,actualizado_por) "
            "VALUES (%s,'Cuenta fallida','Agrupa por tema','Resume temas',1,1) RETURNING id", (customer_id,),
        ).fetchone()["id"]
        # Un objeto huérfano preexistente hace que fn_crear_y_adjuntar omita esa hoja.
        orphan_name = f"registro_regex_c{identifier}"
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
            superuser.execute(sql.SQL("CREATE TABLE dat.{} (cuenta_id int)").format(sql.Identifier(orphan_name)))
        try:
            with pytest.raises(psycopg.Error):
                api_conn.execute("SELECT ops.fn_aprovisionar_cuenta(%s,1)", (identifier,))
            api_conn.rollback()
            with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
                assert superuser.execute("SELECT count(*) FROM cfg.cuenta WHERE id=%s", (identifier,)).fetchone()[0] == 0
                assert superuser.execute("SELECT count(*) FROM aud.auditoria WHERE modulo='cuentas' AND entidad_id=%s",
                                         (str(identifier),)).fetchone()[0] == 0
                assert superuser.execute(
                    "SELECT count(*) FROM pg_inherits i JOIN pg_class child ON child.oid=i.inhrelid "
                    "WHERE child.relname LIKE %s", (f"%_c{identifier}",),
                ).fetchone()[0] == 0
        finally:
            with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
                superuser.execute(sql.SQL("DROP TABLE dat.{}").format(sql.Identifier(orphan_name)))


def test_cuentas_api_crud_etag_visibilidad_y_particiones(environment):
    client, password, port = environment
    admin = _login(client, password)
    consultor = _login(client, password, "consultor@syntra.local")
    assert client.get("/api/v1/cuentas").status_code == 401
    assert client.post("/api/v1/cuentas", headers=consultor, json={}).status_code == 403
    customer = client.post("/api/v1/clientes", headers=admin, json={"nombre": "Cliente para API de Cuentas"})
    assert customer.status_code == 201, customer.text
    customer_id = customer.json()["id"]
    body = {"cliente_id": customer_id, "nombre": "Cuenta de análisis", "taxonomia": "gobierno",
            "prompt_clusters": "Agrupa cada registro por tema principal.",
            "prompt_temas": "Resume los temas finales de la conversación.", "estado_id": 1}
    created = client.post("/api/v1/cuentas", headers=admin, json=body)
    assert created.status_code == 201, created.text
    account_id = created.json()["id"]
    assert created.headers["ETag"] and created.headers["Location"].endswith(f"/{account_id}")
    assert client.post("/api/v1/cuentas", headers=admin, json=body).status_code == 409
    listed_accounts = client.get("/api/v1/cuentas", headers=admin).json()
    assert listed_accounts["total_aprox"] >= 1
    assert any(row["id"] == account_id for row in listed_accounts["items"])
    assert client.get(f"/api/v1/cuentas/{account_id}", headers=consultor).status_code == 404
    assert all(row["id"] != account_id for row in client.get("/api/v1/cuentas", headers=consultor).json()["items"])
    assert client.get(f"/api/v1/cuentas/{account_id}/historial", headers=consultor).status_code == 403
    detail = client.get(f"/api/v1/cuentas/{account_id}", headers=admin)
    assert detail.status_code == 200
    assert client.patch(f"/api/v1/cuentas/{account_id}", headers=admin,
                        json={"nombre": "Cuenta editada"}).status_code == 428
    assert client.patch(f"/api/v1/cuentas/{account_id}", headers={**admin, "If-Match": '"old"'},
                        json={"nombre": "Cuenta editada"}).status_code == 412
    changed = client.patch(f"/api/v1/cuentas/{account_id}",
                           headers={**admin, "If-Match": detail.headers["ETag"]},
                           json={"nombre": "Cuenta editada"})
    assert changed.status_code == 200 and changed.json()["nombre"] == "Cuenta editada"
    assert client.post(f"/api/v1/cuentas/{account_id}/estado", headers=admin,
                       json={"estado_id": 0}).json()["estado_id"] == 0
    assert client.delete(f"/api/v1/cuentas/{account_id}", headers=admin).status_code == 204
    assert client.post(f"/api/v1/cuentas/{account_id}/restaurar", headers=admin).json()["estado_id"] == 0
    history = client.get(f"/api/v1/cuentas/{account_id}/historial", headers=admin)
    assert history.status_code == 200 and len(history.json()) == 5
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        assert superuser.execute(
            "SELECT count(*) FROM pg_inherits i JOIN pg_class child ON child.oid=i.inhrelid "
            "WHERE child.relname LIKE %s", (f"%_c{account_id}",),
        ).fetchone()[0] == 12
        assert superuser.execute("SELECT count(*) FROM aud.auditoria WHERE modulo='cuentas' AND entidad_id=%s",
                                 (str(account_id),)).fetchone()[0] == 5


def test_grupos_api_vinculos_visibilidad_y_auditoria(environment):
    client, password, port = environment
    admin = _login(client, password)
    consultor = _login(client, password, "consultor@syntra.local")
    catalog = client.get("/api/v1/usuarios", headers=admin)
    assert catalog.status_code == 200 and {1, 2} <= {row["id"] for row in catalog.json()["items"]}
    assert "password_hash" not in catalog.text and "version_token" not in catalog.text
    assert client.get("/api/v1/usuarios", headers=consultor).status_code == 403
    assert client.get("/api/v1/grupos").status_code == 401
    assert client.post("/api/v1/grupos", headers=consultor, json={"nombre": "No autorizado"}).status_code == 403
    customer = client.post("/api/v1/clientes", headers=admin, json={"nombre": "Cliente de grupo aislado"})
    assert customer.status_code == 201, customer.text
    account = client.post("/api/v1/cuentas", headers=admin, json={
        "cliente_id": customer.json()["id"], "nombre": "Cuenta de grupo aislado",
        "prompt_clusters": "Agrupa cada registro por tema principal.",
        "prompt_temas": "Resume los temas finales de la conversación.",
    })
    assert account.status_code == 201, account.text
    account_id = account.json()["id"]
    assert client.get(f"/api/v1/cuentas/{account_id}", headers=consultor).status_code == 404
    created = client.post("/api/v1/grupos", headers=admin, json={
        "nombre": " Equipo  Norte ", "descripcion": "  Seguimiento   regional ",
        "usuario_ids": [2], "cuenta_ids": [account_id], "estado_id": 1,
    })
    assert created.status_code == 201, created.text
    group_id = created.json()["id"]
    assert created.json()["nombre"] == "Equipo Norte"
    assert created.json()["usuario_ids"] == [2] and created.json()["cuenta_ids"] == [account_id]
    assert client.get(f"/api/v1/cuentas/{account_id}", headers=consultor).status_code == 200
    assert client.post("/api/v1/grupos", headers=admin, json={"nombre": "equipo norte"}).status_code == 409
    assert client.post("/api/v1/grupos", headers=admin, json={"nombre": "Otro", "cuenta_ids": [999999]}).status_code == 400
    assert client.patch(f"/api/v1/grupos/{group_id}", headers=admin, json={"descripcion": "Cambio"}).status_code == 428
    assert client.patch(f"/api/v1/grupos/{group_id}", headers={**admin, "If-Match": '"viejo"'},
                        json={"descripcion": "Cambio"}).status_code == 412
    detail = client.get(f"/api/v1/grupos/{group_id}", headers=admin)
    edited = client.patch(f"/api/v1/grupos/{group_id}", headers={**admin, "If-Match": detail.headers["ETag"]},
                          json={"descripcion": "Cambio", "cuenta_ids": []})
    assert edited.status_code == 200, edited.text
    assert edited.json()["usuario_ids"] == [2] and edited.json()["cuenta_ids"] == []
    assert client.get(f"/api/v1/cuentas/{account_id}", headers=consultor).status_code == 404
    edited_again = client.patch(f"/api/v1/grupos/{group_id}",
                                headers={**admin, "If-Match": edited.headers["ETag"]},
                                json={"cuenta_ids": [account_id]})
    assert edited_again.status_code == 200, edited_again.text
    assert client.post(f"/api/v1/grupos/{group_id}/estado", headers=admin,
                       json={"estado_id": 0}).json()["estado_id"] == 0
    assert client.get(f"/api/v1/cuentas/{account_id}", headers=consultor).status_code == 404
    assert client.delete(f"/api/v1/grupos/{group_id}", headers=admin).status_code == 204
    restored = client.post(f"/api/v1/grupos/{group_id}/restaurar", headers=admin)
    assert restored.status_code == 200 and restored.json()["estado_id"] == 0
    assert restored.json()["cuenta_ids"] == [account_id]
    assert client.post(f"/api/v1/grupos/{group_id}/estado", headers=admin,
                       json={"estado_id": 1}).json()["estado_id"] == 1
    assert client.get(f"/api/v1/cuentas/{account_id}", headers=consultor).status_code == 200
    history = client.get(f"/api/v1/grupos/{group_id}/historial", headers=admin)
    assert history.status_code == 200 and len(history.json()) == 7
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
        assert superuser.execute("SELECT count(*) FROM seg.grupo_usuario WHERE grupo_id=%s", (group_id,)).fetchone()[0] == 1
        assert superuser.execute("SELECT count(*) FROM seg.grupo_cuenta WHERE grupo_id=%s", (group_id,)).fetchone()[0] == 1
        assert superuser.execute("SELECT count(*) FROM aud.auditoria WHERE modulo='grupos' AND entidad_id=%s",
                                 (str(group_id),)).fetchone()[0] == 7


def test_auditoria_listado_paginado_y_clave_completa(environment):
    client, password, _port = environment
    admin = _login(client, password)
    consultor = _login(client, password, "consultor@syntra.local")
    assert client.get("/api/v1/auditoria").status_code == 401
    assert client.get("/api/v1/auditoria", headers=consultor).status_code == 403
    created = client.post("/api/v1/clientes", headers=admin, json={"nombre": "Cliente auditable"})
    assert created.status_code == 201, created.text
    identifier = created.json()["id"]
    changed = client.patch(f"/api/v1/clientes/{identifier}", headers={**admin, "If-Match": created.headers["ETag"]},
                           json={"sector": "Investigación"})
    assert changed.status_code == 200, changed.text
    params = {"modulo": "clientes", "entidad_tipo": "Cliente", "entidad_id": str(identifier), "limite": 1}
    first = client.get("/api/v1/auditoria", headers=admin, params=params)
    assert first.status_code == 200, first.text
    assert len(first.json()["items"]) == 1 and first.json()["siguiente"]
    second = client.get("/api/v1/auditoria", headers=admin,
                        params={**params, "cursor": first.json()["siguiente"]})
    assert second.status_code == 200 and len(second.json()["items"]) == 1
    assert second.json()["siguiente"] is None
    assert client.get("/api/v1/auditoria", headers=admin,
                      params={**params, "modulo": "cuentas", "cursor": first.json()["siguiente"]}).status_code == 400
    assert client.get("/api/v1/auditoria", headers=admin,
                      params={"desde": "2026-10-22T00:00:00Z", "hasta": "2026-10-20T00:00:00Z"}).status_code == 400
    event = first.json()["items"][0]
    detail = client.get(f"/api/v1/auditoria/{event['id']}", headers=admin, params={"fecha": event["fecha"]})
    assert detail.status_code == 200 and detail.json()["id"] == event["id"]
    assert client.get(f"/api/v1/auditoria/{event['id']}", headers=admin).status_code == 400
    assert client.get(f"/api/v1/auditoria/{event['id']}", headers=admin,
                      params={"fecha": "2020-01-01T00:00:00Z"}).status_code == 404
    period = {"desde": "2026-01-01T00:00:00Z", "hasta": "2027-01-01T00:00:00Z"}
    assert client.get("/api/v1/auditoria/resumen", params=period).status_code == 401
    assert client.get("/api/v1/auditoria/resumen", headers=consultor, params=period).status_code == 403
    summary = client.get("/api/v1/auditoria/resumen", headers=admin, params=period)
    assert summary.status_code == 200, summary.text
    assert sum(day["total"] for day in summary.json()["por_dia"]) >= 2
    assert next(row["total"] for row in summary.json()["por_modulo"] if row["modulo"] == "clientes") >= 2
    assert client.get("/api/v1/auditoria/resumen", headers=admin,
                      params={"desde": "2027-01-01T00:00:00Z", "hasta": "2026-01-01T00:00:00Z"}).status_code == 400
    assert client.get("/api/v1/auditoria/resumen", headers=admin,
                      params={"desde": "2026-01-01", "hasta": "2027-01-01"}).status_code == 400


def test_auditoria_exportacion_worker_xlsx_descarga_y_privilegios(environment, monkeypatch):
    from app.api.routers import auditoria as audit_router

    client, password, port = environment
    admin = _login(client, password)
    consultor = _login(client, password, "consultor@syntra.local")
    with tempfile.TemporaryDirectory(prefix="syntra_reports_") as temp:
        root = Path(temp) / "reports"
        monkeypatch.setattr(audit_router, "settings", replace(audit_router.settings, report_root=root))
        name = "Cliente para exportación real"
        created = client.post("/api/v1/clientes", headers=admin, json={"nombre": name})
        assert created.status_code == 201, created.text
        history = client.get(f"/api/v1/clientes/{created.json()['id']}/historial", headers=admin)
        assert history.status_code == 200 and any(event["accion"] == "crear" for event in history.json())
        assert client.get(f"/api/v1/clientes/{created.json()['id']}/historial", headers=consultor).status_code == 403
        filters = {"modulo": "clientes", "entidad_tipo": "Cliente", "entidad_id": str(created.json()["id"])}
        idem = "auditoria-export-trial-0001"
        accepted = client.post("/api/v1/auditoria/exportar", headers={**admin, "Idempotency-Key": idem}, json=filters)
        assert accepted.status_code == 202, accepted.text
        process_id = accepted.json()["proceso_id"]
        # Un evento posterior al instante de corte no debe aparecer en ese XLSX.
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
            superuser.execute(
                "INSERT INTO aud.auditoria(fecha,modulo,accion,entidad_tipo,entidad_id,detalle) "
                "VALUES (now()+interval '1 minute','clientes','editar','Cliente',%s,'evento posterior al corte')",
                (str(created.json()["id"]),),
            )
        assert accepted.headers["Location"] == f"/api/v1/procesos/{process_id}"
        same = client.post("/api/v1/auditoria/exportar", headers={**admin, "Idempotency-Key": idem}, json=filters)
        assert same.status_code == 202 and same.json()["proceso_id"] == process_id
        conflict = client.post("/api/v1/auditoria/exportar", headers={**admin, "Idempotency-Key": idem},
                               json={**filters, "accion": "editar"})
        assert conflict.status_code == 409
        assert client.get(f"/api/v1/procesos/{process_id}", headers=consultor).status_code == 404
        assert client.get(f"/api/v1/auditoria/exportaciones/{process_id}/descarga", headers=admin).status_code == 409
        code = """
import json
from app.config import WorkerSettings
from app.reports import ReportQueue
queue = ReportQueue(WorkerSettings.from_environment())
task = queue.claim('trial-report-worker')
assert task is not None
assert queue.claim('trial-report-worker-2') is None
queue.process(task, 'trial-report-worker')
from pathlib import Path
from time import time
import os
old = time() - 9 * 86400
for name in ('auditoria-999.tmp', 'auditoria-999.xlsx'):
    file = queue.root / name
    file.write_bytes(b'trial')
    os.utime(file, (old, old))
removed = queue.cleanup()
print(json.dumps({'proceso_id': task.process_id, 'limpieza': removed}))
"""
        env = {**os.environ, "SYNTRA_WORKER_CONCURRENCY": "2", "SYNTRA_WORKER_QUEUE": "q.reportes",
               "SYNTRA_WORKER_DB_USER": "syntra_report_worker_login", "SYNTRA_DB_HOST": "127.0.0.1",
               "SYNTRA_DB_PORT": str(port), "SYNTRA_DB_NAME": "syntra_trial", "SYNTRA_REPORT_ROOT": str(root)}
        result = subprocess.run([sys.executable, "-c", code], cwd=Path(__file__).resolve().parents[2] / "worker",
                                env=env, capture_output=True, text=True, timeout=90, check=False)
        assert result.returncode == 0, result.stderr[-1200:]
        assert json.loads(result.stdout)["proceso_id"] == process_id
        assert json.loads(result.stdout)["limpieza"] == [1, 1]
        assert not (root / "auditoria-999.tmp").exists() and not (root / "auditoria-999.xlsx").exists()
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="syntra_report_worker_login") as limited:
            assert limited.execute("SELECT has_table_privilege(current_user,'aud.auditoria','SELECT')").fetchone()[0]
            assert not limited.execute("SELECT has_schema_privilege(current_user,'seg','USAGE')").fetchone()[0]
            assert not limited.execute("SELECT has_schema_privilege(current_user,'dat','CREATE')").fetchone()[0]
        status = client.get(f"/api/v1/procesos/{process_id}", headers=admin)
        assert status.status_code == 200 and status.json()["estado"] == "completado", status.text
        assert status.json()["total"] == 1
        assert client.post("/api/v1/auditoria/exportar", headers={**admin, "Idempotency-Key": idem},
                           json=filters).json()["estado"] == "completado"
        link = client.get(f"/api/v1/auditoria/exportaciones/{process_id}/descarga", headers=admin)
        assert link.status_code == 200, link.text
        assert "?" not in link.json()["url"] and "token" not in link.json()
        cookie_header = link.headers["Set-Cookie"].lower()
        assert "secure" in cookie_header and "httponly" in cookie_header and "samesite=strict" in cookie_header
        assert client.get(link.json()["url"], headers=consultor).status_code == 403
        file = client.get(link.json()["url"], headers=admin)
        assert file.status_code == 200 and file.content.startswith(b"PK"), file.text[:300]
        sheet = load_workbook(BytesIO(file.content), read_only=True).worksheets[0]
        rows = list(sheet.values)
        assert rows[0][0:3] == ("Fecha", "ID", "Usuario")
        assert len(rows) == 2 and name in rows[1]
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
            superuser.execute("UPDATE prc.proceso SET finalizado_en=now()-interval '8 days' WHERE id=%s", (process_id,))
        assert client.get(f"/api/v1/auditoria/exportaciones/{process_id}/descarga", headers=admin).status_code == 410

        retry = client.post("/api/v1/auditoria/exportar", headers={**admin, "Idempotency-Key": "auditoria-export-trial-retry-0001"}, json=filters)
        assert retry.status_code == 202
        retry_id = retry.json()["proceso_id"]
        failure_code = code.replace("queue.process(task, 'trial-report-worker')",
                                    "queue._write = lambda *args: (_ for _ in ()).throw(OSError('disco'))\nqueue.process(task, 'trial-report-worker')")
        failed = subprocess.run([sys.executable, "-c", failure_code], cwd=Path(__file__).resolve().parents[2] / "worker",
                                env=env, capture_output=True, text=True, timeout=90, check=False)
        assert failed.returncode == 0, failed.stderr[-1200:]
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
            task = superuser.execute("SELECT estado,intentos,error_codigo FROM prc.proceso_tarea WHERE proceso_id=%s",
                                     (retry_id,)).fetchone()
            assert task == ("en_cola", 1, "exportacion_fallida")
            assert not (root / f"auditoria-{retry_id}.xlsx").exists()
            superuser.execute("UPDATE prc.proceso_tarea SET disponible_en=now() WHERE proceso_id=%s", (retry_id,))
        recovered = subprocess.run([sys.executable, "-c", code], cwd=Path(__file__).resolve().parents[2] / "worker",
                                   env=env, capture_output=True, text=True, timeout=90, check=False)
        assert recovered.returncode == 0 and json.loads(recovered.stdout)["proceso_id"] == retry_id, recovered.stderr[-1200:]
        assert client.get(f"/api/v1/procesos/{retry_id}", headers=admin).json()["estado"] == "completado"

        interrupted = client.post("/api/v1/auditoria/exportar",
                                  headers={**admin, "Idempotency-Key": "auditoria-export-trial-crash-0001"}, json=filters)
        assert interrupted.status_code == 202
        interrupted_id = interrupted.json()["proceso_id"]
        claim_code = """from app.config import WorkerSettings
from app.reports import ReportQueue
queue = ReportQueue(WorkerSettings.from_environment())
task = queue.claim('trial-crashed-worker')
assert task is not None
print(task.process_id)
"""
        claimed = subprocess.run([sys.executable, "-c", claim_code], cwd=Path(__file__).resolve().parents[2] / "worker",
                                 env=env, capture_output=True, text=True, timeout=90, check=False)
        assert claimed.returncode == 0 and int(claimed.stdout.strip()) == interrupted_id, claimed.stderr[-1200:]
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
            superuser.execute("UPDATE prc.proceso_tarea SET latido=now()-interval '16 minutes' WHERE proceso_id=%s",
                              (interrupted_id,))
        rescue_code = code.replace("task = queue.claim('trial-report-worker')",
                                   "assert queue.recover() == 1\ntask = queue.claim('trial-report-worker')")
        rescued = subprocess.run([sys.executable, "-c", rescue_code], cwd=Path(__file__).resolve().parents[2] / "worker",
                                 env=env, capture_output=True, text=True, timeout=90, check=False)
        assert rescued.returncode == 0 and json.loads(rescued.stdout)["proceso_id"] == interrupted_id, rescued.stderr[-1200:]
        assert client.get(f"/api/v1/procesos/{interrupted_id}", headers=admin).json()["estado"] == "completado"

        terminal = client.post("/api/v1/auditoria/exportar",
                               headers={**admin, "Idempotency-Key": "auditoria-export-trial-terminal-0001"}, json=filters)
        assert terminal.status_code == 202
        terminal_id = terminal.json()["proceso_id"]
        for attempt in range(1, 4):
            failed = subprocess.run([sys.executable, "-c", failure_code],
                                    cwd=Path(__file__).resolve().parents[2] / "worker",
                                    env=env, capture_output=True, text=True, timeout=90, check=False)
            assert failed.returncode == 0 and json.loads(failed.stdout)["proceso_id"] == terminal_id, failed.stderr[-1200:]
            with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
                task_state = superuser.execute(
                    "SELECT estado,intentos,error_codigo FROM prc.proceso_tarea WHERE proceso_id=%s", (terminal_id,)
                ).fetchone()
                assert task_state == (("error" if attempt == 3 else "en_cola"), attempt, "exportacion_fallida")
                if attempt < 3:
                    superuser.execute("UPDATE prc.proceso_tarea SET disponible_en=now() WHERE proceso_id=%s", (terminal_id,))
        assert client.get(f"/api/v1/procesos/{terminal_id}", headers=admin).json()["estado"] == "error"
        assert client.get(f"/api/v1/auditoria/exportaciones/{terminal_id}/descarga", headers=admin).status_code == 409
        assert not (root / f"auditoria-{terminal_id}.xlsx").exists()

        lost = client.post("/api/v1/auditoria/exportar",
                           headers={**admin, "Idempotency-Key": "auditoria-export-trial-lost-lease-0001"}, json=filters)
        assert lost.status_code == 202
        lost_id = lost.json()["proceso_id"]
        stage_code = """import json,time
from app.config import WorkerSettings
from app.reports import ReportQueue
queue = ReportQueue(WorkerSettings.from_environment())
task = queue.claim('trial-old-worker')
assert task is not None
request, timeout = queue.load(task)
count, staged = queue.render(task, request, time.monotonic()+timeout/1000)
print(json.dumps({'id':task.id,'process_id':task.process_id,'attempts':task.attempts,
                  'count':count,'staged':str(staged)}))
"""
        staged_result = subprocess.run([sys.executable, "-c", stage_code],
                                       cwd=Path(__file__).resolve().parents[2] / "worker",
                                       env=env, capture_output=True, text=True, timeout=90, check=False)
        assert staged_result.returncode == 0, staged_result.stderr[-1200:]
        old = json.loads(staged_result.stdout)
        assert old["process_id"] == lost_id and Path(old["staged"]).is_file()
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as superuser:
            superuser.execute("UPDATE prc.proceso_tarea SET latido=now()-interval '16 minutes' WHERE proceso_id=%s",
                              (lost_id,))
        reclaim_code = """import json
from app.config import WorkerSettings
from app.reports import ReportQueue
queue = ReportQueue(WorkerSettings.from_environment())
assert queue.recover() == 1
task = queue.claim('trial-new-worker')
assert task is not None
print(json.dumps({'id':task.id,'process_id':task.process_id,'attempts':task.attempts}))
"""
        reclaimed_result = subprocess.run([sys.executable, "-c", reclaim_code],
                                          cwd=Path(__file__).resolve().parents[2] / "worker",
                                          env=env, capture_output=True, text=True, timeout=90, check=False)
        assert reclaimed_result.returncode == 0, reclaimed_result.stderr[-1200:]
        new = json.loads(reclaimed_result.stdout)
        assert new["id"] == old["id"] and new["attempts"] == 2
        stale_code = f"""from pathlib import Path
from app.config import WorkerSettings
from app.domain.tasks import ClaimedTask
from app.reports import ReportQueue
queue = ReportQueue(WorkerSettings.from_environment())
assert queue.publish(ClaimedTask({old['id']},{lost_id},1,{{}}),'trial-old-worker',
                     {old['count']},Path({old['staged']!r})) is False
"""
        stale_result = subprocess.run([sys.executable, "-c", stale_code],
                                      cwd=Path(__file__).resolve().parents[2] / "worker",
                                      env=env, capture_output=True, text=True, timeout=90, check=False)
        assert stale_result.returncode == 0, stale_result.stderr[-1200:]
        assert not Path(old["staged"]).exists() and not (root / f"auditoria-{lost_id}.xlsx").exists()
        new_code = f"""from app.config import WorkerSettings
from app.domain.tasks import ClaimedTask
from app.reports import ReportQueue
queue = ReportQueue(WorkerSettings.from_environment())
queue.process(ClaimedTask({new['id']},{lost_id},2,{{}}),'trial-new-worker')
"""
        new_result = subprocess.run([sys.executable, "-c", new_code],
                                    cwd=Path(__file__).resolve().parents[2] / "worker",
                                    env=env, capture_output=True, text=True, timeout=90, check=False)
        assert new_result.returncode == 0, new_result.stderr[-1200:]
        assert client.get(f"/api/v1/procesos/{lost_id}", headers=admin).json()["estado"] == "completado"
        assert (root / f"auditoria-{lost_id}.xlsx").is_file()


@pytest.fixture
def monitor_environment(environment):
    """Base aislada para que la prueba de colas no dependa del orden de la suite."""
    client, password, port = environment
    database_name = f"syntra_monitor_{secrets.token_hex(8)}"
    root = Path(__file__).resolve().parents[2]
    original_override = app.dependency_overrides.get(production_connection)

    with psycopg.connect(host="127.0.0.1", port=port, dbname="postgres", user="postgres", autocommit=True) as admin:
        admin.execute(sql.SQL("CREATE DATABASE {} TEMPLATE syntra_trial").format(sql.Identifier(database_name)))

    try:
        with psycopg.connect(host="127.0.0.1", port=port, dbname=database_name, user="postgres") as db:
            db.execute("TRUNCATE TABLE prc.proceso CASCADE")
            db.execute("TRUNCATE TABLE ops.worker")

        def monitor_connection():
            with psycopg.connect(host="127.0.0.1", port=port, dbname=database_name,
                                 user="syntra_api_login", row_factory=psycopg.rows.dict_row) as conn:
                yield conn

        app.dependency_overrides[production_connection] = monitor_connection
        yield client, password, port, database_name
    finally:
        if original_override is None:
            app.dependency_overrides.pop(production_connection, None)
        else:
            app.dependency_overrides[production_connection] = original_override
        with psycopg.connect(host="127.0.0.1", port=port, dbname="postgres", user="postgres", autocommit=True) as admin:
            admin.execute(sql.SQL("DROP DATABASE IF EXISTS {}").format(sql.Identifier(database_name)))


def test_operations_monitor_real_queries(monitor_environment):
    """Consulta API y SQL real en el PostgreSQL descartable, nunca en syntra_dev."""
    client, password, port, database_name = monitor_environment
    admin_headers = _login(client, password)
    owner_id = client.get("/api/v1/auth/yo", headers=admin_headers).json()["usuario"]["id"]
    worker_id = f"monitor-test:{secrets.token_hex(8)}"
    active_old_worker_id = f"monitor-test-old-active:{secrets.token_hex(8)}"
    ingest_worker_id = f"monitor-test-ingest:{secrets.token_hex(8)}"

    with psycopg.connect(host="127.0.0.1", port=port, dbname=database_name, user="postgres") as db:
        # Concesiones temporales y mínimas en el clúster de prueba para representar
        # el acceso que necesitará syntra_api_login al habilitar el monitor local.
        db.execute("GRANT SELECT (id,servidor,colas,iniciado_en,latido,tareas_en_curso) ON ops.worker TO syntra_api_login")
        db.execute("GRANT SELECT (id,proceso_id,clave,tipo,estado,intentos,ok,errores,error_codigo,iniciado_en,finalizado_en,cola,disponible_en) ON prc.proceso_tarea TO syntra_api_login")
        db.execute("GRANT SELECT (medicion_id,modo,duracion_ms,sla_ms) ON prc.proceso TO syntra_api_login")

    no_workers = client.get("/api/v1/ops/workers", headers=admin_headers)
    assert no_workers.status_code == 200 and no_workers.json() == []
    empty_queues = client.get("/api/v1/ops/colas", headers=admin_headers)
    assert empty_queues.status_code == 200
    assert all(row["en_cola"] == row["en_ejecucion"] == row["con_error"] == 0
               for row in empty_queues.json())

    with psycopg.connect(host="127.0.0.1", port=port, dbname=database_name, user="postgres") as db:
        db.execute(
            "INSERT INTO ops.worker(id,servidor,colas,pid,iniciado_en,latido,tareas_en_curso) "
            "VALUES (%s,'monitor-test',ARRAY['q.reportes'],321,now()-interval '2 minutes',now()-interval '60 seconds',0)",
            (worker_id,),
        )
        db.execute(
            "INSERT INTO ops.worker(id,servidor,colas,pid,iniciado_en,latido,tareas_en_curso) "
            "VALUES (%s,'monitor-test',ARRAY['q.reportes'],322,now()-interval '8 days',now(),0)",
            (active_old_worker_id,),
        )
        db.execute(
            "INSERT INTO ops.worker(id,servidor,colas,pid,iniciado_en,latido,tareas_en_curso) "
            "VALUES (%s,'monitor-test',ARRAY['q.ingesta'],323,now()-interval '8 days',now(),0)",
            (ingest_worker_id,),
        )
        for process_type, kind, queue, task_type, state, code in (
            ("exportacion", "auditoria", "q.reportes", "auditoria_exportar", "en_cola", None),
            ("exportacion", "auditoria", "q.reportes", "auditoria_exportar", "en_ejecucion", None),
            ("exportacion", "auditoria", "q.reportes", "auditoria_exportar", "completado", None),
            ("exportacion", "auditoria", "q.reportes", "auditoria_exportar", "error", "exportacion_fallida"),
            ("correo_acceso", None, "q.notificaciones", "correo_acceso", "en_cola", None),
            ("carga_archivo", None, "q.ingesta", "carga_archivo", "en_cola", None),
        ):
            options = json.dumps({"kind": kind} if kind else {})
            process_id = db.execute(
                "INSERT INTO prc.proceso(tipo,modo,estado,opciones,creado_por) "
                "VALUES (%s,'manual',%s,%s::jsonb,%s) RETURNING id",
                (process_type, state, options, owner_id),
            ).fetchone()[0]
            if state == "completado":
                db.execute(
                    "UPDATE prc.proceso SET iniciado_en=now()-interval '25 seconds',finalizado_en=now(),sla_ms=60000,procesados=3 WHERE id=%s",
                    (process_id,),
                )
            db.execute(
                "INSERT INTO prc.proceso_tarea(proceso_id,clave,tipo,cola,estado,intentos,error_codigo,error_mensaje) "
                "VALUES (%s,'monitor',%s,%s,%s,1,%s,'mensaje técnico secreto')",
                (process_id, task_type, queue, state, code),
            )
        db.execute(
            """INSERT INTO prc.proceso(tipo,modo,estado,opciones,creado_por,creado_en)
               SELECT 'exportacion','manual','completado','{"kind":"auditoria"}'::jsonb,%s,now()-interval '1 hour'
                 FROM generate_series(1,27)""",
            (owner_id,),
        )
        db.execute(
            """INSERT INTO prc.proceso(tipo,modo,estado,opciones,creado_por,creado_en)
               VALUES ('exportacion','manual','completado','{"kind":"auditoria"}'::jsonb,%s,now()-interval '8 days')""",
            (owner_id,),
        )

    response = client.get("/api/v1/ops/colas", headers=admin_headers)
    assert response.status_code == 200, response.text
    queue_rows = {row["cola"]: row for row in response.json()}
    assert queue_rows["q.notificaciones"]["en_cola"] >= 1
    assert queue_rows["q.reportes"]["en_cola"] >= 1
    assert queue_rows["q.reportes"]["en_ejecucion"] >= 1
    assert queue_rows["q.reportes"]["con_error"] >= 1
    assert queue_rows["q.reportes"]["workers"] == 1
    assert queue_rows["q.ingesta"]["en_cola"] == 1
    assert queue_rows["q.ingesta"]["workers"] == 1

    workers = client.get("/api/v1/ops/workers", headers=admin_headers)
    assert workers.status_code == 200, workers.text
    stale_worker = next(row for row in workers.json() if row["id"] == worker_id)
    assert stale_worker["vivo"] is False
    assert stale_worker["tareas_en_curso"] == 0
    old_but_active_worker = next(row for row in workers.json() if row["id"] == active_old_worker_id)
    assert old_but_active_worker["vivo"] is True
    assert next(row for row in workers.json() if row["id"] == ingest_worker_id)["vivo"] is True
    with psycopg.connect(host="127.0.0.1", port=port, dbname=database_name, user="postgres") as db:
        db.execute("UPDATE ops.worker SET latido=now()-interval '60 seconds' WHERE id=%s", (ingest_worker_id,))
    ingest_after_heartbeat = client.get("/api/v1/ops/colas", headers=admin_headers)
    assert ingest_after_heartbeat.status_code == 200
    ingest_queue = next(row for row in ingest_after_heartbeat.json() if row["cola"] == "q.ingesta")
    assert ingest_queue["en_cola"] == 1 and ingest_queue["workers"] == 0

    processes = client.get("/api/v1/procesos?limite=25", headers=admin_headers)
    assert processes.status_code == 200, processes.text
    assert len(processes.json()) == 25
    assert processes.headers.get("X-Next-Cursor")
    next_page = client.get("/api/v1/procesos?limite=25&cursor=" + processes.headers["X-Next-Cursor"], headers=admin_headers)
    assert next_page.status_code == 200, next_page.text
    assert len(next_page.json()) == 7
    all_page_rows = processes.json() + next_page.json()
    process_ids = [row["id"] for row in all_page_rows]
    assert len(process_ids) == len(set(process_ids)) == 32
    creation_order = [(datetime.fromisoformat(row["creado_en"].replace("Z", "+00:00")), row["id"]) for row in all_page_rows]
    assert creation_order == sorted(creation_order, reverse=True)
    assert all(datetime.now(timezone.utc) - timedelta(days=7) <= datetime.fromisoformat(row["creado_en"].replace("Z", "+00:00")) <= datetime.now(timezone.utc) for row in all_page_rows)
    monitor_processes = [row for row in processes.json() if row["tipo"] in {"correo_acceso", "exportacion"}]
    assert {row["tipo"] for row in monitor_processes} == {"correo_acceso", "exportacion"}
    statistics = client.get("/api/v1/procesos/estadisticas", headers=admin_headers)
    assert statistics.status_code == 200, statistics.text
    summary = statistics.json()
    assert summary["en_ejecucion_o_en_cola"] >= 2
    assert summary["completados_24h"] >= 1
    assert summary["registros_procesados_24h"] >= 3
    assert 24000 <= summary["duracion_promedio_ms"] <= 26000
    assert summary["sla_cumplimiento_pct"] == 100.0
    assert summary["ultimas_duraciones"]
    failed = next(row for row in monitor_processes if row["tipo"] == "exportacion" and row["estado"] == "error")
    task_response = client.get(f"/api/v1/procesos/{failed['id']}/tareas", headers=admin_headers)
    assert task_response.status_code == 200, task_response.text
    assert task_response.json()[0]["error_codigo"] == "exportacion_fallida"
    assert "mensaje técnico secreto" not in task_response.text

    # La matriz aprobada asigna procesos:ver al Consultor; se revoca solo en esta
    # base desechable para verificar el caso negativo real de autorización.
    with psycopg.connect(host="127.0.0.1", port=port, dbname=database_name, user="postgres") as db:
        db.execute("DELETE FROM seg.rol_permiso WHERE rol_id=4 AND modulo='procesos' AND accion='ver'")
    consultant_headers = _login(client, password, email="consultor@syntra.local")
    assert client.get("/api/v1/ops/colas", headers=consultant_headers).status_code == 403


def test_measurement_creation_and_worker_preparation_are_persistent_and_retryable(environment):
    """Alta, configuración, tarea durable, worker y reintentos en PostgreSQL temporal."""
    client, password, port = environment
    admin_headers = _login(client, password)
    actor_id = client.get("/api/v1/auth/yo", headers=admin_headers).json()["usuario"]["id"]
    suffix = secrets.token_hex(6)
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
        db.execute("INSERT INTO cat.origen(id,codigo,nombre) VALUES (5,'manual','Manual de prueba') ON CONFLICT DO NOTHING")
        client_id = db.execute(
            "INSERT INTO cfg.cliente(nombre,creado_por,actualizado_por) VALUES (%s,%s,%s) RETURNING id",
            (f"Cliente de integración {suffix}", actor_id, actor_id),
        ).fetchone()[0]
        account_id = db.execute(
            "INSERT INTO cfg.cuenta(cliente_id,nombre,prompt_clusters,prompt_temas,creado_por,actualizado_por) "
            "VALUES (%s,%s,'prompt de clusters','prompt de temas',%s,%s) RETURNING id",
            (client_id, f"Cuenta de integración {suffix}", actor_id, actor_id),
        ).fetchone()[0]

    def create_measurement(name: str):
        result = client.post("/api/v1/mediciones", headers=admin_headers, json={
            "cuenta_id": account_id,
            "nombre": name,
            "fuentes": [{"origen": "manual", "activo": True}],
            "prompts": [{"tipo": "cluster", "origen": "propio", "contenido": "Prompt congelado de prueba"}],
            "cruces": [{"nombre": f"Cruce de integración {suffix}", "items": ["#Syntra"]}],
        })
        assert result.status_code == 201, result.text
        return result.json()

    created = create_measurement(f"Medición de integración {suffix}")
    measurement_id = created["id"]
    process_id = created["proceso_preparacion_id"]
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
        durable = db.execute(
            "SELECT p.estado,t.estado,t.cola,t.tipo,t.payload->>'medicion_id',t.payload->>'actor_id' "
            "FROM prc.proceso p JOIN prc.proceso_tarea t ON t.proceso_id=p.id "
            "WHERE p.id=%s AND p.medicion_id=%s", (process_id, measurement_id),
        ).fetchone()
        assert durable == ("en_cola", "en_cola", "q.mantenimiento", "preparar_medicion",
                           str(measurement_id), str(actor_id))
        assert db.execute("SELECT contenido FROM med.medicion_prompt WHERE medicion_id=%s", (measurement_id,)).fetchone() == ("Prompt congelado de prueba",)
        assert db.execute(
            "SELECT count(*) FROM med.medicion_cruce mc JOIN cfg.cruce_item i ON i.cruce_id=mc.cruce_id "
            "WHERE mc.medicion_id=%s AND i.item='#Syntra'", (measurement_id,),
        ).fetchone()[0] == 1
        assert db.execute(
            "SELECT count(*) FROM aud.auditoria WHERE entidad_tipo='Medición' AND entidad_id=%s",
            (str(measurement_id),),
        ).fetchone()[0] == 1
        assert db.execute(
            "SELECT has_function_privilege('syntra_measurement_worker_login',"
            "'ops.fn_preparar_medicion(integer,integer,integer[])','EXECUTE'),"
            "has_function_privilege('syntra_measurement_worker_login',"
            "'ops.fn_preparar_medicion_segura(bigint,text,integer,integer,integer,integer)','EXECUTE'),"
            "has_table_privilege('syntra_measurement_worker_login','ops.worker','DELETE')"
        ).fetchone() == (False, True, False)

    root = Path(__file__).resolve().parents[2]
    env = os.environ.copy()
    env.update({
        "SYNTRA_WORKER_QUEUE": "q.mantenimiento",
        "SYNTRA_WORKER_CONCURRENCY": "1",
        "SYNTRA_DB_HOST": "127.0.0.1",
        "SYNTRA_DB_PORT": str(port),
        "SYNTRA_DB_NAME": "syntra_trial",
        "SYNTRA_WORKER_DB_USER": "syntra_measurement_worker_login",
        "SYNTRA_WORKER_DB_PASSFILE": str(Path(tempfile.gettempdir()) / f"syntra-unset-passfile-{suffix}"),
    })
    worker_id = f"trial:{suffix}:q.mantenimiento:prep"
    prepare_code = f"""from app.config import WorkerSettings
from app.infrastructure.measurement_queue import MeasurementQueue
from app.infrastructure.measurement_worker_registry import PostgresMeasurementWorkerRegistry
settings = WorkerSettings.from_environment()
registry = PostgresMeasurementWorkerRegistry(settings)
queue = MeasurementQueue(settings)
worker_id = {worker_id!r}
registry.register(worker_id, 'trial', 123)
registry.heartbeat(worker_id, 1)
try:
    task = queue.claim(worker_id)
    assert task is not None
    queue.prepare(task, worker_id)
finally:
    registry.unregister(worker_id)
"""
    prepared = subprocess.run([sys.executable, "-c", prepare_code], cwd=root / "worker", env=env,
                              capture_output=True, text=True, timeout=90, check=False)
    assert prepared.returncode == 0, prepared.stderr[-1800:]
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
        assert db.execute(
            "SELECT p.estado,t.estado,p.progreso FROM prc.proceso p "
            "JOIN prc.proceso_tarea t ON t.proceso_id=p.id WHERE p.id=%s", (process_id,),
        ).fetchone() == ("completado", "completado", 100)
        assert db.execute(
            "SELECT count(*) FROM ops.fn_tablas_dat() tab "
            "WHERE to_regclass('dat.' || quote_ident(tab || '_c' || %s::text || '_m' || %s::text)) IS NOT NULL",
            (account_id, measurement_id),
        ).fetchone()[0] == 12
        assert db.execute("SELECT count(*) FROM ops.worker WHERE id=%s", (worker_id,)).fetchone()[0] == 0

    retry = create_measurement(f"Medición reintentos {suffix}")
    retry_process = retry["proceso_preparacion_id"]
    for attempt, approved_ms in ((1, 30_000), (2, 120_000), (3, None)):
        retry_after_query = "    assert row[0] == 'error'" if approved_ms is None else (
            f"    assert row[0] == 'en_cola' and {approved_ms - 1000} <= row[1] <= {approved_ms}\n"
            f"    conn.execute('UPDATE prc.proceso_tarea SET disponible_en=now() WHERE proceso_id=%s', ({retry_process},))"
        )
        exercise_retry = f"""import psycopg
from app.config import WorkerSettings
from app.infrastructure.measurement_queue import MeasurementQueue
settings = WorkerSettings.from_environment()
queue = MeasurementQueue(settings)
worker_id = {worker_id!r}
task = queue.claim(worker_id)
assert task is not None and task.attempts == {attempt}
queue.fail(task, worker_id, 'preparacion_fallida')
with psycopg.connect(host='127.0.0.1', port={port}, dbname='syntra_trial',
                     user='syntra_measurement_worker_login',
                     passfile=str(settings.database_passfile)) as conn:
    row = conn.execute('SELECT estado, ceil(extract(epoch FROM (disponible_en-now()))*1000)::int '
                       'FROM prc.proceso_tarea WHERE proceso_id=%s', ({retry_process},)).fetchone()
{retry_after_query}
"""
        failed = subprocess.run([sys.executable, "-c", exercise_retry], cwd=root / "worker", env=env,
                                capture_output=True, text=True, timeout=45, check=False)
        assert failed.returncode == 0, failed.stderr[-1800:]
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
        retry_result = db.execute(
            "SELECT p.estado,t.estado,t.intentos,t.error_codigo FROM prc.proceso p "
            "JOIN prc.proceso_tarea t ON t.proceso_id=p.id WHERE p.id=%s", (retry_process,),
        ).fetchone()
        assert retry_result == ("error", "error", 3, "preparacion_fallida")

    # Recuperación tras caída: rescatar una tarea con lease vencido, respetar
    # el primer backoff persistido y permitir que un worker nuevo la complete.
    abandoned = create_measurement(f"Medición recuperación {suffix}")
    abandoned_process = abandoned["proceso_preparacion_id"]
    claim_code = f"""from app.config import WorkerSettings
from app.infrastructure.measurement_queue import MeasurementQueue
settings = WorkerSettings.from_environment()
task = MeasurementQueue(settings).claim({worker_id!r})
assert task is not None and task.process_id == {abandoned_process}
"""
    claimed = subprocess.run([sys.executable, "-c", claim_code], cwd=root / "worker", env=env,
                             capture_output=True, text=True, timeout=45, check=False)
    assert claimed.returncode == 0, claimed.stderr[-1500:]
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
        db.execute("UPDATE prc.proceso_tarea SET latido=now()-interval '10 minutes' WHERE proceso_id=%s",
                   (abandoned_process,))
    recover_code = f"""from app.config import WorkerSettings
from app.infrastructure.measurement_queue import MeasurementQueue
settings = WorkerSettings.from_environment()
assert MeasurementQueue(settings).recover() == 1
"""
    recovered = subprocess.run([sys.executable, "-c", recover_code], cwd=root / "worker", env=env,
                               capture_output=True, text=True, timeout=45, check=False)
    assert recovered.returncode == 0, recovered.stderr[-1500:]
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
        row = db.execute(
            "SELECT p.estado,t.estado,t.intentos,t.error_codigo, "
            "ceil(extract(epoch FROM (t.disponible_en-now()))*1000)::int "
            "FROM prc.proceso p JOIN prc.proceso_tarea t ON t.proceso_id=p.id WHERE p.id=%s",
            (abandoned_process,),
        ).fetchone()
        assert row[:4] == ("en_cola", "en_cola", 1, "worker_interrumpido")
        assert 29_000 <= row[4] <= 30_000
        db.execute("UPDATE prc.proceso_tarea SET disponible_en=now() WHERE proceso_id=%s", (abandoned_process,))
    recovery_worker_id = worker_id.replace(":prep", ":recovered")
    recovery_code = f"""from app.config import WorkerSettings
from app.infrastructure.measurement_queue import MeasurementQueue
from app.infrastructure.measurement_worker_registry import PostgresMeasurementWorkerRegistry
settings = WorkerSettings.from_environment()
registry = PostgresMeasurementWorkerRegistry(settings)
queue = MeasurementQueue(settings)
worker_id = {recovery_worker_id!r}
registry.register(worker_id, 'trial', 124)
registry.heartbeat(worker_id, 1)
try:
    task = queue.claim(worker_id)
    assert task is not None and task.attempts == 2
    queue.prepare(task, worker_id)
finally:
    registry.unregister(worker_id)
"""
    completed_after_recovery = subprocess.run([sys.executable, "-c", recovery_code], cwd=root / "worker", env=env,
                                               capture_output=True, text=True, timeout=90, check=False)
    assert completed_after_recovery.returncode == 0, completed_after_recovery.stderr[-1800:]


    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
        assert db.execute("SELECT estado FROM prc.proceso WHERE id=%s", (abandoned_process,)).fetchone()[0] == "completado"

    # Dos consumidores concurrentes deben reclamar procesos distintos y completar ambos.
    concurrent = [create_measurement(f"Medición concurrente {suffix}-{i}") for i in range(2)]
    expected_processes = {item["proceso_preparacion_id"] for item in concurrent}
    parallel_codes = []
    for index in range(2):
        parallel_worker_id = worker_id.replace(":prep", f":parallel{index}")
        parallel_codes.append(f"""from app.config import WorkerSettings
from app.infrastructure.measurement_queue import MeasurementQueue
from app.infrastructure.measurement_worker_registry import PostgresMeasurementWorkerRegistry
settings = WorkerSettings.from_environment()
registry = PostgresMeasurementWorkerRegistry(settings)
queue = MeasurementQueue(settings)
worker_id = {parallel_worker_id!r}
registry.register(worker_id, 'trial', {130 + index})
registry.heartbeat(worker_id, 1)
try:
    task = queue.claim(worker_id)
    assert task is not None
    print(task.process_id, flush=True)
    queue.prepare(task, worker_id)
finally:
    registry.unregister(worker_id)
""")
    processes = [subprocess.Popen([sys.executable, "-c", code], cwd=root / "worker", env=env,
                                  stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
                 for code in parallel_codes]
    parallel_results = [process.communicate(timeout=90) for process in processes]
    assert all(process.returncode == 0 for process in processes), parallel_results
    claimed_processes = {int(stdout.strip()) for stdout, _stderr in parallel_results}
    assert claimed_processes == expected_processes
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
        states = db.execute("SELECT estado FROM prc.proceso WHERE id=ANY(%s)",
                             (list(expected_processes),)).fetchall()
        assert states == [("completado",), ("completado",)]


def test_measurement_step2_configuration_round_trips_and_patches_with_etag(environment):
    client, password, port = environment
    admin_headers = _login(client, password)
    actor_id = client.get("/api/v1/auth/yo", headers=admin_headers).json()["usuario"]["id"]
    suffix = secrets.token_hex(6)
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
        for origin_id, code, label in ((5, "manual", "Manual"), (6, "meltwater", "Meltwater"), (7, "archivo", "Archivo")):
            db.execute("INSERT INTO cat.origen(id,codigo,nombre) VALUES (%s,%s,%s) ON CONFLICT (codigo) DO NOTHING",
                       (origin_id, code, label))
        client_id = db.execute(
            "INSERT INTO cfg.cliente(nombre,creado_por,actualizado_por) VALUES (%s,%s,%s) RETURNING id",
            (f"Cliente paso 2 {suffix}", actor_id, actor_id),
        ).fetchone()[0]
        account_id = db.execute(
            "INSERT INTO cfg.cuenta(cliente_id,nombre,prompt_clusters,prompt_temas,creado_por,actualizado_por) "
            "VALUES (%s,%s,'prompt clusters','prompt temas',%s,%s) RETURNING id",
            (client_id, f"Cuenta paso 2 {suffix}", actor_id, actor_id),
        ).fetchone()[0]
    created = client.post("/api/v1/mediciones", headers=admin_headers, json={
        "cuenta_id": account_id, "nombre": f"Medición paso 2 {suffix}",
        "fuentes": [
            {"origen": "manual", "activo": True, "config": {},
             "al_cargar": {"incorporar": False, "relevancia": True, "sentimiento": False, "clusters": True, "bots": True}},
            {"origen": "meltwater", "activo": False, "config": {"busqueda_id": "MW-TEMP"},
             "al_cargar": {"incorporar": True, "relevancia": False, "sentimiento": True, "clusters": False, "bots": False}},
        ], "auto_publicar": False,
    })
    assert created.status_code == 201, created.text
    measurement_id = created.json()["id"]
    detail = client.get(f"/api/v1/mediciones/{measurement_id}", headers=admin_headers)
    assert detail.status_code == 200, detail.text
    assert detail.headers.get("ETag")
    before = detail.json()
    source_before = {item["origen"]: item for item in before["fuentes"]}
    assert source_before["manual"]["al_cargar"] == {
        "incorporar": False, "relevancia": True, "sentimiento": False, "clusters": True, "bots": True,
    }
    assert source_before["meltwater"]["config"]["busqueda_id"] == "MW-TEMP"
    assert before["programacion_meltwater"]["activa"] is False

    payload = {
        "fuentes": [
            {"origen": "manual", "activo": True, "config": {},
             "al_cargar": {"incorporar": True, "relevancia": False, "sentimiento": True, "clusters": False, "bots": False}},
            {"origen": "meltwater", "activo": True, "config": {"busqueda_id": "MW-CHANGED"},
             "al_cargar": {"incorporar": False, "relevancia": True, "sentimiento": True, "clusters": True, "bots": False}},
        ],
        "programacion_meltwater": {"activa": True, "inicio": "2026-10-23T00:00:00Z",
                                   "fin": "2026-10-24T00:00:00Z", "frecuencia_min": 30},
        "auto_publicar": True,
    }
    updated = client.patch(f"/api/v1/mediciones/{measurement_id}", headers={**admin_headers, "If-Match": detail.headers["ETag"]}, json=payload)
    assert updated.status_code == 200, updated.text
    assert updated.headers.get("ETag") and updated.headers["ETag"] != detail.headers["ETag"]
    persisted = client.get(f"/api/v1/mediciones/{measurement_id}", headers=admin_headers)
    assert persisted.status_code == 200, persisted.text
    after = persisted.json()
    sources_after = {item["origen"]: item for item in after["fuentes"]}
    assert sources_after["meltwater"]["activo"] is True
    assert sources_after["meltwater"]["config"]["busqueda_id"] == "MW-CHANGED"
    assert sources_after["meltwater"]["al_cargar"]["relevancia"] is True
    assert after["programacion_meltwater"]["activa"] is True
    assert after["programacion_meltwater"]["frecuencia_min"] == 30
    assert after["auto_publicar"] is True
    stale = client.patch(f"/api/v1/mediciones/{measurement_id}", headers={**admin_headers, "If-Match": detail.headers["ETag"]}, json={"auto_publicar": False})
    assert stale.status_code == 412, stale.text
    missing = client.patch(f"/api/v1/mediciones/{measurement_id}", headers=admin_headers, json={"auto_publicar": False})
    assert missing.status_code == 428, missing.text
    assert client.get("/api/v1/mediciones", headers=admin_headers).status_code == 200
    assert client.get("/api/v1/mediciones/999999", headers=admin_headers).status_code == 404


def test_measurement_file_upload_queue_and_tenant_boundaries(environment, monkeypatch):
    """API, archivo privado y cola reales en PostgreSQL descartable."""
    from app.api.routers import measurement_files

    client, password, port = environment
    admin = _login(client, password)
    template = client.get("/api/v1/formatos/talkwalker_simple/plantilla", headers=admin)
    assert template.status_code == 200, template.text
    assert template.json()["nombre"] == "plantilla_talkwalker_simple.xlsx"
    downloaded_template = client.get(template.json()["url"], headers=admin)
    assert downloaded_template.status_code == 200
    assert downloaded_template.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    workbook = load_workbook(BytesIO(downloaded_template.content), read_only=True)
    assert workbook.sheetnames == ["Plantilla"]
    assert next(workbook.active.iter_rows(values_only=True))
    assert client.get("/api/v1/formatos/talkwalker_simple/plantilla?token=invalid",
                      headers=admin).status_code == 403
    actor = client.get("/api/v1/auth/yo", headers=admin).json()["usuario"]["id"]
    suffix = secrets.token_hex(6)
    with tempfile.TemporaryDirectory(prefix="syntra_upload_trial_") as storage:
        monkeypatch.setattr(measurement_files, "settings",
                            replace(measurement_files.settings, upload_root=Path(storage)))
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
            client_id = db.execute(
                "INSERT INTO cfg.cliente(nombre,creado_por,actualizado_por) VALUES (%s,%s,%s) RETURNING id",
                (f"Cliente archivo {suffix}", actor, actor),
            ).fetchone()[0]
            account_id = db.execute(
                "INSERT INTO cfg.cuenta(cliente_id,nombre,prompt_clusters,prompt_temas,creado_por,actualizado_por) "
                "VALUES (%s,%s,'clusters','temas',%s,%s) RETURNING id",
                (client_id, f"Cuenta archivo {suffix}", actor, actor),
            ).fetchone()[0]
        created = client.post("/api/v1/mediciones", headers=admin, json={
            "cuenta_id": account_id, "nombre": f"Medición archivo {suffix}",
            "fuentes": [{"origen": "archivo", "activo": True}],
        })
        assert created.status_code == 201, created.text
        measurement_id = created.json()["id"]
        data = b"Date,Headline,Opening Text,Hit Sentence,Source,Influencer,Country,Reach,Sentiment,URL\n"
        reserved = client.post(f"/api/v1/mediciones/{measurement_id}/archivos/presignar",
                               headers=admin, json={"nombre": "talkwalker.csv", "tamano_bytes": len(data),
                                                    "formato_id": "talkwalker_simple"})
        assert reserved.status_code == 201, reserved.text
        file_id = reserved.json()["archivo_id"]
        unmapped = client.post(f"/api/v1/mediciones/{measurement_id}/archivos/presignar",
                               headers=admin, json={"nombre": "meltwater.xlsx", "tamano_bytes": 10,
                                                    "formato_id": "meltwater"})
        assert unmapped.status_code == 400, unmapped.text
        uploaded = client.put(f"/api/v1/archivos/{file_id}/contenido", headers=admin, content=data)
        assert uploaded.status_code == 204, uploaded.text
        listed = client.get(f"/api/v1/mediciones/{measurement_id}/archivos", headers=admin)
        assert listed.status_code == 200, listed.text
        assert listed.json()[0]["id"] == file_id
        extra = client.post(f"/api/v1/mediciones/{measurement_id}/archivos/presignar",
                            headers=admin, json={"nombre": "pendiente.csv", "tamano_bytes": 8,
                                                 "formato_id": "talkwalker_simple"})
        assert extra.status_code == 201, extra.text
        first_page = client.get(f"/api/v1/mediciones/{measurement_id}/archivos?limite=1", headers=admin)
        assert first_page.status_code == 200 and len(first_page.json()) == 1
        # La reserva sin contenido no es una carga del historial.
        assert first_page.json()[0]["id"] == file_id
        assert first_page.json()[0]["subido_por_nombre"]
        assert not first_page.headers.get("X-Next-Cursor")
        assert client.get(f"/api/v1/mediciones/{measurement_id}/archivos?cursor=invalid!",
                          headers=admin).status_code == 400
        wrong_hash = client.post(f"/api/v1/mediciones/{measurement_id}/archivos/presignar",
                                 headers=admin, json={"nombre": "hash.csv", "tamano_bytes": len(data),
                                                      "formato_id": "talkwalker_simple", "sha256": "0" * 64})
        assert wrong_hash.status_code == 201, wrong_hash.text
        assert client.put(f"/api/v1/archivos/{wrong_hash.json()['archivo_id']}/contenido",
                          headers=admin, content=data).status_code == 400
        hashed_data = data + b'\n'
        declared_hash = __import__('hashlib').sha256(hashed_data).hexdigest()
        right_hash = client.post(f"/api/v1/mediciones/{measurement_id}/archivos/presignar",
                                 headers=admin, json={"nombre": "verified.csv", "tamano_bytes": len(hashed_data),
                                                      "formato_id": "talkwalker_simple", "sha256": declared_hash})
        assert right_hash.status_code == 201, right_hash.text
        assert client.put(f"/api/v1/archivos/{right_hash.json()['archivo_id']}/contenido",
                          headers=admin, content=hashed_data).status_code == 204
        first_page = client.get(f"/api/v1/mediciones/{measurement_id}/archivos?limite=1", headers=admin)
        assert first_page.status_code == 200 and [item["id"] for item in first_page.json()] == [right_hash.json()["archivo_id"]]
        cursor = first_page.headers.get("X-Next-Cursor")
        assert cursor
        second_page = client.get(f"/api/v1/mediciones/{measurement_id}/archivos?limite=1&cursor={cursor}",
                                 headers=admin)
        assert second_page.status_code == 200 and [item["id"] for item in second_page.json()] == [file_id]
        assert not second_page.headers.get("X-Next-Cursor")
        assert client.get(f"/api/v1/archivos/{file_id}/contenido", headers=admin).status_code == 403
        link = client.get(f"/api/v1/archivos/{file_id}/descargar", headers=admin)
        assert link.status_code == 200, link.text
        downloaded = client.get(link.json()["url"], headers=admin)
        assert downloaded.status_code == 200 and downloaded.content == data
        key = "test-file-ingest-" + suffix + "-123456789"
        queued = client.post(f"/api/v1/mediciones/{measurement_id}/cargas/archivos",
                             headers={**admin, "Idempotency-Key": key},
                             json={"archivo_ids": [file_id]})
        assert queued.status_code == 202, queued.text
        repeated = client.post(f"/api/v1/mediciones/{measurement_id}/cargas/archivos",
                               headers={**admin, "Idempotency-Key": key},
                               json={"archivo_ids": [file_id]})
        assert repeated.status_code == 202 and repeated.json() == queued.json()
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
            assert db.execute("SELECT count(*) FROM prc.proceso_tarea WHERE proceso_id=%s AND cola='q.ingesta'",
                              (queued.json()["proceso_id"],)).fetchone()[0] == 1
            row = db.execute("SELECT sha256,s3_bucket FROM med.archivo WHERE id=%s", (file_id,)).fetchone()
            assert row == (__import__("hashlib").sha256(data).digest(), "local")
            other_client = db.execute(
                "INSERT INTO cfg.cliente(nombre,creado_por,actualizado_por) VALUES (%s,%s,%s) RETURNING id",
                (f"Otro cliente archivo {suffix}", actor, actor),
            ).fetchone()[0]
            other_account = db.execute(
                "INSERT INTO cfg.cuenta(cliente_id,nombre,prompt_clusters,prompt_temas,creado_por,actualizado_por) "
                "VALUES (%s,%s,'clusters','temas',%s,%s) RETURNING id",
                (other_client, f"Otra cuenta archivo {suffix}", actor, actor),
            ).fetchone()[0]
            role_id = db.execute(
                "INSERT INTO seg.rol(nombre,ver_todas,estado_id) VALUES (%s,false,1) RETURNING id",
                (f"Archivo acotado {suffix}",),
            ).fetchone()[0]
            db.execute("INSERT INTO seg.rol_permiso(rol_id,modulo,accion) VALUES "
                       "(%s,'mediciones','ver'),(%s,'mediciones','ejecutar')", (role_id, role_id))
            other_user = db.execute(
                "INSERT INTO seg.usuario(nombre,email,password_hash,rol_id) VALUES (%s,%s,%s,%s) RETURNING id",
                ("Usuario otra cuenta", f"archivo-{suffix}@syntra.local", password_hasher.hash(password), role_id),
            ).fetchone()[0]
            group_id = db.execute("INSERT INTO seg.grupo(nombre) VALUES (%s) RETURNING id",
                                  (f"Grupo archivo {suffix}",)).fetchone()[0]
            db.execute("INSERT INTO seg.grupo_usuario(grupo_id,usuario_id) VALUES (%s,%s)", (group_id, other_user))
            db.execute("INSERT INTO seg.grupo_cuenta(grupo_id,cuenta_id) VALUES (%s,%s)", (group_id, other_account))
        scoped = _login(client, password, f"archivo-{suffix}@syntra.local")
        assert client.get(f"/api/v1/mediciones/{measurement_id}/archivos", headers=scoped).status_code == 404
        assert client.get(f"/api/v1/archivos/{file_id}/descargar", headers=scoped).status_code == 404
        assert client.post(f"/api/v1/mediciones/{measurement_id}/archivos/presignar",
                           headers=scoped, json={"nombre": "ajeno.csv", "tamano_bytes": len(data),
                                                 "formato_id": "talkwalker_simple"}).status_code == 404
        assert client.put(f"/api/v1/archivos/{file_id}/contenido", headers=scoped, content=data).status_code == 404
        assert client.post(f"/api/v1/mediciones/{measurement_id}/cargas/archivos",
                           headers={**scoped, "Idempotency-Key": "tenant-denied-123456789"},
                           json={"archivo_ids": [file_id]}).status_code == 404
        # La fixture comparte el clúster descartable entre casos: cerrar esta
        # tarea deliberadamente no consumida antes de probar otro consumidor.
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
            db.execute("UPDATE prc.proceso_tarea SET estado='error',finalizado_en=now() "
                       "WHERE proceso_id=%s AND estado='en_cola'", (queued.json()["proceso_id"],))
            db.execute("UPDATE prc.proceso SET estado='error',finalizado_en=now() WHERE id=%s",
                       (queued.json()["proceso_id"],))
            db.execute("UPDATE prc.proceso_tarea SET estado='error',finalizado_en=now() "
                       "WHERE proceso_id IN (SELECT id FROM prc.proceso WHERE medicion_id=%s "
                       "AND tipo='preparar_medicion') AND estado='en_cola'", (measurement_id,))
            db.execute("UPDATE prc.proceso SET estado='error',finalizado_en=now() "
                       "WHERE medicion_id=%s AND tipo='preparar_medicion' AND estado='en_cola'",
                       (measurement_id,))


def _close_previous_measurement_tasks(port: int) -> None:
    """Aisla el clúster compartido de pruebas; nunca se usa en syntra_dev."""
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
        processes = [row[0] for row in db.execute(
            "SELECT DISTINCT proceso_id FROM prc.proceso_tarea "
            "WHERE cola IN ('q.mantenimiento','q.ingesta') AND estado='en_cola'"
        )]
        if processes:
            db.execute("UPDATE prc.proceso_tarea SET estado='error',finalizado_en=now() "
                       "WHERE proceso_id=ANY(%s) AND estado='en_cola'", (processes,))
            db.execute("UPDATE prc.proceso SET estado='error',finalizado_en=now() "
                       "WHERE id=ANY(%s) AND estado='en_cola'", (processes,))


def test_ingest_worker_processes_real_csv_with_checkpoint_and_rejections(environment, monkeypatch):
    """Recorrido API -> archivo local -> q.ingesta -> dat en clúster descartable."""
    from app.api.routers import measurement_files

    client, password, port = environment
    _close_previous_measurement_tasks(port)
    admin = _login(client, password)
    actor = client.get("/api/v1/auth/yo", headers=admin).json()["usuario"]["id"]
    suffix = secrets.token_hex(6)
    with tempfile.TemporaryDirectory(prefix="syntra_ingest_trial_") as storage:
        monkeypatch.setattr(measurement_files, "settings",
                            replace(measurement_files.settings, upload_root=Path(storage)))
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
            client_id = db.execute("INSERT INTO cfg.cliente(nombre) VALUES (%s) RETURNING id",
                                   (f"Cliente ingesta {suffix}",)).fetchone()[0]
            account_id = db.execute("INSERT INTO cfg.cuenta(cliente_id,nombre,prompt_clusters,prompt_temas) "
                                    "VALUES (%s,%s,'clusters','temas') RETURNING id",
                                    (client_id, f"Cuenta ingesta {suffix}")).fetchone()[0]
        created = client.post("/api/v1/mediciones", headers=admin, json={
            "cuenta_id": account_id, "nombre": f"Medición ingesta {suffix}",
            "fuentes": [{"origen": "archivo", "activo": True}],
        })
        assert created.status_code == 201, created.text
        measurement_id = created.json()["id"]
        data = ("url,published,content,source_type\n"
                 'https://example.test/a,07/10/26 12:46:00,Texto A,"SOCIALMEDIA,SOCIALMEDIA_TWITTER"\n'
                 'https://example.test/a,07/10/26 12:46:00,Texto A,"SOCIALMEDIA,SOCIALMEDIA_TWITTER"\n'
                 'https://example.test/old,26/10/07 12:46:00,Texto viejo,"SOCIALMEDIA,SOCIALMEDIA_TWITTER"\n').encode()
        reserved = client.post(f"/api/v1/mediciones/{measurement_id}/archivos/presignar",
                               headers=admin, json={"nombre": "talkwalker.csv", "tamano_bytes": len(data),
                                                    "formato_id": "talkwalker_simple"})
        assert reserved.status_code == 201, reserved.text
        file_id = reserved.json()["archivo_id"]
        assert client.put(f"/api/v1/archivos/{file_id}/contenido", headers=admin, content=data).status_code == 204
        queued = client.post(f"/api/v1/mediciones/{measurement_id}/cargas/archivos",
                             headers={**admin, "Idempotency-Key": "ingest-real-" + suffix + "-12345678"},
                             json={"archivo_ids": [file_id], "al_incorporar": {"incorporar": True, "relevancia": True}})
        assert queued.status_code == 202, queued.text
        process_id = queued.json()["proceso_id"]
        prepare_code = """
from app.config import WorkerSettings
from app.infrastructure.measurement_queue import MeasurementQueue
queue = MeasurementQueue(WorkerSettings.from_environment())
task = queue.claim('trial:q.mantenimiento:one')
assert task is not None
queue.process(task, 'trial:q.mantenimiento:one')
"""
        prepare_env = {**os.environ, "SYNTRA_WORKER_CONCURRENCY": "1",
                       "SYNTRA_WORKER_QUEUE": "q.mantenimiento", "SYNTRA_DB_HOST": "127.0.0.1",
                       "SYNTRA_DB_PORT": str(port), "SYNTRA_DB_NAME": "syntra_trial",
                       "SYNTRA_WORKER_DB_USER": "syntra_measurement_worker_login"}
        prepared = subprocess.run([sys.executable, "-c", prepare_code],
                                  cwd=Path(__file__).resolve().parents[2] / "worker", env=prepare_env,
                                  text=True, capture_output=True, timeout=90, check=False)
        assert prepared.returncode == 0, prepared.stderr[-1200:]
        worker_code = """
from app.config import WorkerSettings
from app.infrastructure.ingest_queue import IngestQueue
queue = IngestQueue(WorkerSettings.from_environment())
task = queue.claim('trial:q.ingesta:one')
assert task is not None
queue.process(task, 'trial:q.ingesta:one')
queue._notify_terminal(task.process_id)
"""
        env = {**os.environ, "SYNTRA_WORKER_CONCURRENCY": "1", "SYNTRA_WORKER_QUEUE": "q.ingesta",
               "SYNTRA_DB_HOST": "127.0.0.1", "SYNTRA_DB_PORT": str(port),
               "SYNTRA_DB_NAME": "syntra_trial", "SYNTRA_WORKER_DB_USER": "syntra_ingest_worker_login",
               "SYNTRA_UPLOAD_ROOT": storage}
        result = subprocess.run([sys.executable, "-c", worker_code],
                                cwd=Path(__file__).resolve().parents[2] / "worker", env=env,
                                text=True, capture_output=True, timeout=90, check=False)
        assert result.returncode == 0, result.stderr[-1200:]
        notices = client.get('/api/v1/notificaciones', headers=admin)
        assert notices.status_code == 200, notices.text
        notice = next(item for item in notices.json()['items']
                      if item['enlace'] == f'/mediciones/{measurement_id}?tab=cargas')
        assert notices.json()['no_leidas'] >= 1
        assert notice['titulo'] == 'Carga con advertencias'
        assert notice['tipo'] == 'warn' and not notice['leida']
        with psycopg.connect(host='127.0.0.1', port=port, dbname='syntra_trial', user='postgres') as db:
            assert db.execute('SELECT count(*) FROM aud.notificacion WHERE proceso_id=%s',
                              (process_id,)).fetchone()[0] == 1
        consultor_notice = _login(client, password, 'consultor@syntra.local')
        assert all(item['id'] != notice['id'] for item in client.get(
            '/api/v1/notificaciones', headers=consultor_notice).json()['items'])
        assert client.post(f"/api/v1/notificaciones/{notice['id']}/leer",
                           headers=consultor_notice).status_code == 404
        assert client.post(f"/api/v1/notificaciones/{notice['id']}/leer",
                           headers=admin).status_code == 204
        assert not client.get('/api/v1/notificaciones?solo_no_leidas=true',
                              headers=admin).json()['items']
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
            task = db.execute("SELECT estado,intentos,payload,error_codigo FROM prc.proceso_tarea "
                              "WHERE proceso_id=%s", (process_id,)).fetchone()
            assert task[0] == "completado", (task, result.stderr[-1200:])
            assert task[1] == 1 and task[2]["offset"] == 3 and task[3] is None
            assert db.execute("SELECT estado FROM prc.proceso WHERE id=%s", (process_id,)).fetchone()[0] == "completado"
            assert db.execute("SELECT estado,filas_leidas,filas_validas,filas_rechazadas FROM med.archivo "
                              "WHERE id=%s", (file_id,)).fetchone() == ("procesado", 3, 2, 1)
            assert db.execute("SELECT count(*) FROM dat.registro WHERE medicion_id=%s", (measurement_id,)).fetchone()[0] == 1
            assert db.execute("SELECT count(*) FROM dat.registro_estado WHERE medicion_id=%s", (measurement_id,)).fetchone()[0] == 1
            initial_state = db.execute(
                "SELECT incorporado,incorporado_en,st_relevancia,relevancia FROM dat.registro_estado "
                "WHERE medicion_id=%s", (measurement_id,),
            ).fetchone()
            assert initial_state[0] is True and initial_state[1] is not None
            assert initial_state[2:] == (0, 0)  # evaluación pendiente, incorporación ya confirmada
            assert db.execute("SELECT n_leidos,n_nuevos,n_duplicados,n_rechazados FROM dat.ejecucion "
                              "WHERE tarea_id IN (SELECT id FROM prc.proceso_tarea WHERE proceso_id=%s)",
                              (process_id,)).fetchone() == (3, 1, 1, 1)
            assert db.execute("SELECT count(*) FROM dat.registro_evento WHERE proceso_id=%s", (process_id,)).fetchone()[0] == 1
            rejected_detail = db.execute(
                "SELECT detalle FROM dat.registro_evento WHERE proceso_id=%s", (process_id,),
            ).fetchone()[0]
            assert rejected_detail["fila"] == 3
            assert rejected_detail["original"]["published"] == "26/10/07 12:46:00"

        lots = client.get(f"/api/v1/mediciones/{measurement_id}/lotes?limite=1", headers=admin)
        assert lots.status_code == 200, lots.text
        assert len(lots.json()) == 1
        lot = lots.json()[0]
        assert (lot["n_leidos"], lot["n_nuevos"], lot["n_duplicados"], lot["n_rechazados"]) == (3, 1, 1, 1)
        assert lot["estado"] == "completado" and lot["proceso_id"] == process_id
        assert lot["iniciado_en"] and lot["finalizado_en"] and lot["duracion_ms"] is not None
        formats = client.get("/api/v1/formatos", headers=admin)
        assert formats.status_code == 200, formats.text
        assert any(item["id"] == "talkwalker_simple" and item["columnas"] for item in formats.json())
        summary = client.get(f"/api/v1/mediciones/{measurement_id}/resumen", headers=admin)
        assert summary.status_code == 200, summary.text
        archived_source = next(item for item in summary.json()["por_fuente"] if item["origen"] == "archivo")
        assert (archived_source["cargados"], archived_source["duplicados"]) == (1, 1)
        assert client.get(f"/api/v1/mediciones/{measurement_id}/lotes?cursor=malformado",
                          headers=admin).status_code == 400
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
            origin_id = db.execute("SELECT id FROM cat.origen WHERE codigo='archivo'").fetchone()[0]
            for index in range(3):
                db.execute("INSERT INTO dat.ejecucion(cuenta_id,medicion_id,origen_id,post,detalle,n_duplicados) "
                           "VALUES (%s,%s,%s,'{}',%s,%s)",
                           (account_id, measurement_id, origin_id, f"Lote de paginación {index}", index))
        summary_after = client.get(f"/api/v1/mediciones/{measurement_id}/resumen", headers=admin)
        assert summary_after.status_code == 200, summary_after.text
        previous_loads = next(item for item in summary_after.json()["por_fuente"] if item["origen"] == "archivo")
        assert (previous_loads["cargados"], previous_loads["duplicados"]) == (1, 4)
        first_page = client.get(f"/api/v1/mediciones/{measurement_id}/lotes?limite=2", headers=admin)
        assert first_page.status_code == 200 and len(first_page.json()) == 2
        marker = first_page.headers.get("X-Next-Cursor")
        assert marker
        second_page = client.get(f"/api/v1/mediciones/{measurement_id}/lotes?limite=2&cursor={marker}",
                                 headers=admin)
        assert second_page.status_code == 200 and len(second_page.json()) == 2
        assert not second_page.headers.get("X-Next-Cursor")
        assert len({item["id"] for item in first_page.json() + second_page.json()}) == 4
        assert client.get(f"/api/v1/mediciones/{measurement_id}/lotes?origen=inexistente",
                          headers=admin).json() == []
        consultor = _login(client, password, "consultor@syntra.local")
        assert client.get(f"/api/v1/mediciones/{measurement_id}/lotes",
                          headers=consultor).status_code in (403, 404)
        assert client.get(f"/api/v1/mediciones/{measurement_id}/resumen",
                          headers=consultor).status_code in (403, 404)
        same = client.post(f"/api/v1/mediciones/{measurement_id}/archivos/presignar",
                           headers=admin, json={"nombre": "copia.csv", "tamano_bytes": len(data),
                                                "formato_id": "talkwalker_simple"})
        assert same.status_code == 201, same.text
        duplicate_upload = client.put(f"/api/v1/archivos/{same.json()['archivo_id']}/contenido",
                                      headers=admin, content=data)
        assert duplicate_upload.status_code == 409, duplicate_upload.text
        assert duplicate_upload.json()["detail"] == {
            "codigo": "archivo_duplicado", "archivo_id": file_id,
            "proceso_id": process_id, "estado_archivo": "procesado", "estado_proceso": "completado"}
        duplicate_reserve = client.post(f"/api/v1/mediciones/{measurement_id}/archivos/presignar",
                                        headers=admin, json={"nombre": "copia-por-huella.csv",
                                                             "tamano_bytes": len(data),
                                                             "formato_id": "talkwalker_simple",
                                                             "sha256": hashlib.sha256(data).hexdigest()})
        assert duplicate_reserve.status_code == 409
        assert duplicate_reserve.json()["detail"]["archivo_id"] == file_id
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
            db.execute("UPDATE prc.proceso SET estado='en_ejecucion' WHERE id=%s", (process_id,))
        in_progress = client.post(f"/api/v1/mediciones/{measurement_id}/archivos/presignar",
                                  headers=admin, json={"nombre": "igual-en-curso.csv",
                                                       "tamano_bytes": len(data),
                                                       "formato_id": "talkwalker_simple",
                                                       "sha256": hashlib.sha256(data).hexdigest()})
        assert in_progress.status_code == 409
        assert in_progress.json()["detail"]["estado_proceso"] == "en_ejecucion"
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
            db.execute("UPDATE prc.proceso SET estado='completado' WHERE id=%s", (process_id,))

        # Una huella MD5 heredada se consulta sin reescribirla. La prueba usa
        # exclusivamente el clúster descartable y comprueba la protección de
        # contenido humano, métricas e histórico.
        legacy_hash = hashlib.md5(b"https://example.test/a").digest()
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
            record = db.execute(
                "SELECT cuenta_id,medicion_id,periodo,registro_id,url_hash FROM dat.registro "
                "WHERE medicion_id=%s", (measurement_id,),
            ).fetchone()
            account_scope, _, record_period, record_id, original_hash = record
            assert len(original_hash) == 32
            assert original_hash == db.execute(
                "SELECT ops.fn_hash_registro(%s,NULL,NULL,NULL)",
                ("https://example.test/a",),
            ).fetchone()[0]
            db.execute("UPDATE dat.registro_huella SET url_hash=%s WHERE cuenta_id=%s "
                       "AND medicion_id=%s AND url_hash=%s",
                       (legacy_hash, account_scope, measurement_id, original_hash))
            db.execute("UPDATE dat.registro SET url_hash=%s WHERE cuenta_id=%s AND medicion_id=%s "
                       "AND periodo=%s AND registro_id=%s",
                       (legacy_hash, account_scope, measurement_id, record_period, record_id))
            db.execute("UPDATE dat.registro_estado SET relevancia=1,likes=77 WHERE cuenta_id=%s "
                       "AND medicion_id=%s AND periodo=%s AND registro_id=%s",
                       (account_scope, measurement_id, record_period, record_id))

        classify_code = """
import json, os
from datetime import datetime, timezone
from app.config import WorkerSettings
from app.domain.talkwalker import url_identity
from app.infrastructure.ingest_queue import IngestQueue
canonical, fingerprint, legacy = url_identity(os.environ['SYNTRA_TEST_URL'])
queue = IngestQueue(WorkerSettings.from_environment())
row = {'url': os.environ['SYNTRA_TEST_URL'], 'url_canonica': canonical,
       'url_hash': fingerprint, 'huellas_heredadas': legacy,
       'fecha_publicacion': datetime(2026, 10, 7, 18, 46, tzinfo=timezone.utc),
       'red_id': 1, 'texto': os.environ['SYNTRA_TEST_TEXT'], '_fila': 1}
with queue._connect() as conn:
    accepted, duplicates, protected = queue._classify_chunk(
        conn, [row], {'cuenta_id': int(os.environ['SYNTRA_TEST_ACCOUNT']),
                      'medicion_id': int(os.environ['SYNTRA_TEST_MEASUREMENT'])})
    if os.environ.get('SYNTRA_TEST_ROLLBACK') == '1':
        conn.rollback()
print(json.dumps({'new': len(accepted), 'duplicates': duplicates, 'protected': protected}))
"""
        def classify(url: str, text: str, measurement: int = measurement_id,
                     account: int = account_scope, rollback: bool = False) -> dict:
            variables = {**env, "SYNTRA_TEST_URL": url, "SYNTRA_TEST_TEXT": text,
                         "SYNTRA_TEST_ACCOUNT": str(account),
                         "SYNTRA_TEST_MEASUREMENT": str(measurement),
                         "SYNTRA_TEST_ROLLBACK": '1' if rollback else '0'}
            result = subprocess.run([sys.executable, "-c", classify_code],
                                    cwd=Path(__file__).resolve().parents[2] / "worker",
                                    env=variables, text=True, capture_output=True, timeout=30, check=False)
            assert result.returncode == 0, result.stderr[-1200:]
            return json.loads(result.stdout)

        other_measurement = client.post('/api/v1/mediciones', headers=admin, json={
            'cuenta_id': account_id, 'nombre': f'Medición aislada {suffix}',
            'fuentes': [{'origen': 'archivo', 'activo': True}],
        })
        assert other_measurement.status_code == 201, other_measurement.text
        prepared_other = subprocess.run([sys.executable, '-c', prepare_code],
                                        cwd=Path(__file__).resolve().parents[2] / 'worker',
                                        env=prepare_env, text=True, capture_output=True,
                                        timeout=90, check=False)
        assert prepared_other.returncode == 0, prepared_other.stderr[-1200:]
        assert classify('https://example.test/a', 'Otra medición',
                        measurement=other_measurement.json()['id'], rollback=True)['new'] == 1
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
            other_client = db.execute("INSERT INTO cfg.cliente(nombre) VALUES (%s) RETURNING id",
                                      (f"Cliente aislado ingesta {suffix}",)).fetchone()[0]
            other_account = db.execute(
                "INSERT INTO cfg.cuenta(cliente_id,nombre,prompt_clusters,prompt_temas) "
                "VALUES (%s,%s,'clusters','temas') RETURNING id",
                (other_client, f"Cuenta aislada ingesta {suffix}"),
            ).fetchone()[0]
        other_tenant_measurement = client.post('/api/v1/mediciones', headers=admin, json={
            'cuenta_id': other_account, 'nombre': f'Medición otro cliente {suffix}',
            'fuentes': [{'origen': 'archivo', 'activo': True}],
        })
        assert other_tenant_measurement.status_code == 201, other_tenant_measurement.text
        prepared_tenant = subprocess.run([sys.executable, '-c', prepare_code],
                                         cwd=Path(__file__).resolve().parents[2] / 'worker',
                                         env=prepare_env, text=True, capture_output=True,
                                         timeout=90, check=False)
        assert prepared_tenant.returncode == 0, prepared_tenant.stderr[-1200:]
        assert classify('https://example.test/a', 'Otro cliente',
                        measurement=other_tenant_measurement.json()['id'],
                        account=other_account, rollback=True)['new'] == 1
        assert classify("http://example.test/a/", "Texto nuevo") == {
            "new": 0, "duplicates": 1, "protected": []}
        # Dos cargas simultáneas de la misma URL heredada serializan el upsert.
        with ThreadPoolExecutor(max_workers=2) as pool:
            concurrent_results = list(pool.map(
                lambda _: classify("https://example.test/a/", "Texto concurrente"), range(2)))
        assert concurrent_results == [
            {"new": 0, "duplicates": 1, "protected": []},
            {"new": 0, "duplicates": 1, "protected": []},
        ]
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
            assert db.execute("SELECT texto,url_hash FROM dat.registro WHERE cuenta_id=%s "
                              "AND medicion_id=%s AND periodo=%s AND registro_id=%s",
                              (account_scope, measurement_id, record_period, record_id)).fetchone() == (
                                  "Texto concurrente", legacy_hash)
            assert db.execute("SELECT count(*) FROM dat.registro WHERE cuenta_id=%s AND medicion_id=%s",
                              (account_scope, measurement_id)).fetchone()[0] == 1
            assert db.execute("SELECT relevancia,likes FROM dat.registro_estado WHERE cuenta_id=%s "
                              "AND medicion_id=%s AND periodo=%s AND registro_id=%s",
                              (account_scope, measurement_id, record_period, record_id)).fetchone() == (1, 77)
            db.execute("UPDATE dat.registro SET editado=true,texto='Texto manual' "
                       "WHERE cuenta_id=%s AND medicion_id=%s AND periodo=%s AND registro_id=%s",
                       (account_scope, measurement_id, record_period, record_id))
        assert classify("https://example.test/a", "No debe sobrescribir")['protected'] == []
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
            assert db.execute("SELECT texto FROM dat.registro WHERE cuenta_id=%s "
                              "AND medicion_id=%s AND periodo=%s AND registro_id=%s",
                              (account_scope, measurement_id, record_period, record_id)).fetchone()[0] == "Texto manual"
            db.execute("UPDATE dat.registro_huella SET ubicacion='h' WHERE cuenta_id=%s "
                       "AND medicion_id=%s AND url_hash=%s",
                       (account_scope, measurement_id, legacy_hash))
        historical_location = classify("https://example.test/a", "No debe tocar histórico")
        assert historical_location['new'] == 0 and historical_location['duplicates'] == 1
        assert 'histórico' in historical_location['protected'][0]['reason']
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
            db.execute("UPDATE dat.registro_huella SET ubicacion='p',registro_id=%s "
                       "WHERE cuenta_id=%s AND medicion_id=%s AND url_hash=%s",
                       (record_id + 1000000, account_scope, measurement_id, legacy_hash))
        missing_principal = classify("https://example.test/a", "No debe crear otro")
        assert missing_principal['new'] == 0 and missing_principal['duplicates'] == 1
        assert 'sin registro principal' in missing_principal['protected'][0]['reason']
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
            db.execute("UPDATE dat.registro_huella SET registro_id=%s "
                       "WHERE cuenta_id=%s AND medicion_id=%s AND url_hash=%s",
                       (record_id, account_scope, measurement_id, legacy_hash))
            db.execute("INSERT INTO ops.archivado(cuenta_id,medicion_id,periodo,estado) "
                       "VALUES (%s,%s,%s,'archivado')", (account_scope, measurement_id, record_period))
        historical = classify("https://example.test/a/", "No debe tocar histórico")
        assert historical['new'] == 0 and historical['duplicates'] == 1
        assert 'archivado' in historical['protected'][0]['reason']
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
            assert db.execute("SELECT texto FROM dat.registro WHERE cuenta_id=%s "
                              "AND medicion_id=%s AND periodo=%s AND registro_id=%s",
                              (account_scope, measurement_id, record_period, record_id)).fetchone()[0] == "Texto manual"

        # Un archivo que desaparece antes del procesamiento ejercita los retrasos
        # persistidos sin tocar datos ni archivos de syntra_dev.
        retry_data = data.replace(b"example.test/a", b"example.test/retry")
        second = client.post(f"/api/v1/mediciones/{measurement_id}/archivos/presignar",
                             headers=admin, json={"nombre": "retry.csv", "tamano_bytes": len(retry_data),
                                                  "formato_id": "talkwalker_simple"})
        assert second.status_code == 201, second.text
        second_id = second.json()["archivo_id"]
        assert client.put(f"/api/v1/archivos/{second_id}/contenido", headers=admin,
                          content=retry_data).status_code == 204
        retried = client.post(f"/api/v1/mediciones/{measurement_id}/cargas/archivos",
                              headers={**admin, "Idempotency-Key": "ingest-retry-" + suffix + "-12345678"},
                              json={"archivo_ids": [second_id]})
        assert retried.status_code == 202, retried.text
        retry_process_id = retried.json()["proceso_id"]
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
            key = db.execute("SELECT s3_key FROM med.archivo WHERE id=%s", (second_id,)).fetchone()[0]
        (Path(storage) / key).unlink()
        for attempt, expected_delay in ((1, 30), (2, 120), (3, None)):
            failed = subprocess.run([sys.executable, "-c", worker_code],
                                    cwd=Path(__file__).resolve().parents[2] / "worker", env=env,
                                    text=True, capture_output=True, timeout=90, check=False)
            assert failed.returncode == 0, failed.stderr[-1200:]
            with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
                row = db.execute(
                    "SELECT estado,intentos,error_codigo,"
                    "extract(epoch FROM (disponible_en-now())) FROM prc.proceso_tarea "
                    "WHERE proceso_id=%s", (retry_process_id,),
                ).fetchone()
                assert row[1] == attempt and row[2] == "ingesta_fallida"
                if expected_delay is not None:
                    assert row[0] == "en_cola" and expected_delay - 10 <= row[3] <= expected_delay
                    db.execute("UPDATE prc.proceso_tarea SET disponible_en=now() WHERE proceso_id=%s",
                               (retry_process_id,))
                else:
                    assert row[0] == "error"
                    assert db.execute("SELECT estado FROM med.archivo WHERE id=%s", (second_id,)).fetchone()[0] == "error"
                    assert db.execute("SELECT estado FROM prc.proceso WHERE id=%s",
                                      (retry_process_id,)).fetchone()[0] == "error"

        rescue_data = data.replace(b"example.test/a", b"example.test/resume")
        rescue_file = client.post(f"/api/v1/mediciones/{measurement_id}/archivos/presignar",
                                  headers=admin, json={"nombre": "rescue.csv", "tamano_bytes": len(rescue_data),
                                                       "formato_id": "talkwalker_simple"})
        assert rescue_file.status_code == 201, rescue_file.text
        rescue_file_id = rescue_file.json()["archivo_id"]
        assert client.put(f"/api/v1/archivos/{rescue_file_id}/contenido", headers=admin,
                          content=rescue_data).status_code == 204
        rescue = client.post(f"/api/v1/mediciones/{measurement_id}/cargas/archivos",
                             headers={**admin, "Idempotency-Key": "ingest-rescue-" + suffix + "-12345678"},
                             json={"archivo_ids": [rescue_file_id]})
        assert rescue.status_code == 202, rescue.text
        rescue_process = rescue.json()["proceso_id"]
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
            db.execute("UPDATE prc.proceso_tarea SET estado='en_ejecucion',worker='dead:q.ingesta',"
                       "intentos=1,latido=now()-interval '2 minutes' WHERE proceso_id=%s",
                       (rescue_process,))
        recovered = subprocess.run([sys.executable, "-c",
                                    "from app.config import WorkerSettings\n"
                                    "from app.infrastructure.ingest_queue import IngestQueue\n"
                                    "assert IngestQueue(WorkerSettings.from_environment()).recover() == 1\n"],
                                   cwd=Path(__file__).resolve().parents[2] / "worker", env=env,
                                   text=True, capture_output=True, timeout=90, check=False)
        assert recovered.returncode == 0, recovered.stderr[-1200:]
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
            row = db.execute("SELECT estado,intentos,error_codigo,"
                             "extract(epoch FROM (disponible_en-now())) FROM prc.proceso_tarea "
                             "WHERE proceso_id=%s", (rescue_process,)).fetchone()
            assert row[0:3] == ("en_cola", 1, "worker_interrumpido") and 20 <= row[3] <= 30
            db.execute("UPDATE prc.proceso_tarea SET disponible_en=now() WHERE proceso_id=%s",
                       (rescue_process,))
        resumed = subprocess.run([sys.executable, "-c", worker_code],
                                 cwd=Path(__file__).resolve().parents[2] / "worker", env=env,
                                 text=True, capture_output=True, timeout=90, check=False)
        assert resumed.returncode == 0, resumed.stderr[-1200:]
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
            assert db.execute("SELECT estado,intentos FROM prc.proceso_tarea WHERE proceso_id=%s",
                              (rescue_process,)).fetchone() == ("completado", 2)
            assert db.execute("SELECT count(*) FROM dat.registro WHERE medicion_id=%s "
                              "AND url='https://example.test/resume'", (measurement_id,)).fetchone()[0] == 1

        # La preparación fallida cierra las cargas bloqueadas sin reclamarlas.
        blocked_data = data.replace(b"example.test/a", b"example.test/blocked")
        third = client.post(f"/api/v1/mediciones/{measurement_id}/archivos/presignar",
                            headers=admin, json={"nombre": "blocked.csv", "tamano_bytes": len(blocked_data),
                                                 "formato_id": "talkwalker_simple"})
        assert third.status_code == 201, third.text
        third_id = third.json()["archivo_id"]
        assert client.put(f"/api/v1/archivos/{third_id}/contenido", headers=admin,
                          content=blocked_data).status_code == 204
        blocked = client.post(f"/api/v1/mediciones/{measurement_id}/cargas/archivos",
                              headers={**admin, "Idempotency-Key": "ingest-blocked-" + suffix + "-12345678"},
                              json={"archivo_ids": [third_id]})
        assert blocked.status_code == 202, blocked.text
        blocked_process_id = blocked.json()["proceso_id"]
        # Incluso si un payload y un proceso falsificados concuerdan entre sí,
        # la función acotada comprueba la cuenta real de med.medicion.
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
            alien_client = db.execute("INSERT INTO cfg.cliente(nombre) VALUES (%s) RETURNING id",
                                      (f"Cliente ajeno ingesta {suffix}",)).fetchone()[0]
            alien_account = db.execute(
                "INSERT INTO cfg.cuenta(cliente_id,nombre,prompt_clusters,prompt_temas) "
                "VALUES (%s,%s,'clusters','temas') RETURNING id",
                (alien_client, f"Cuenta ajena ingesta {suffix}"),
            ).fetchone()[0]
            blocked_task_id = db.execute("UPDATE prc.proceso_tarea SET estado='en_ejecucion',"
                                         "worker='trial:q.ingesta:forged',intentos=1,"
                                         "payload=jsonb_set(payload,'{cuenta_id}',to_jsonb(%s::integer)) "
                                         "WHERE proceso_id=%s RETURNING id",
                                         (alien_account, blocked_process_id)).fetchone()[0]
            db.execute("UPDATE prc.proceso SET cuenta_id=%s WHERE id=%s",
                       (alien_account, blocked_process_id))
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial",
                             user="syntra_ingest_worker_login") as limited:
            with pytest.raises(psycopg.errors.InsufficientPrivilege):
                limited.execute("SELECT ops.fn_asegurar_periodos_ingesta(%s,%s,%s)",
                                (blocked_task_id, 'trial:q.ingesta:forged', [202610]))
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
            db.execute("UPDATE prc.proceso SET cuenta_id=%s WHERE id=%s",
                       (account_id, blocked_process_id))
            db.execute("UPDATE prc.proceso_tarea SET estado='en_cola',worker=NULL,intentos=0,"
                       "payload=jsonb_set(payload,'{cuenta_id}',to_jsonb(%s::integer)) WHERE id=%s",
                       (account_id, blocked_task_id))
            db.execute("UPDATE prc.proceso SET estado='error' WHERE medicion_id=%s "
                       "AND tipo='preparar_medicion'", (measurement_id,))
        recovery = subprocess.run([sys.executable, "-c",
                                   "from app.config import WorkerSettings\n"
                                   "from app.infrastructure.ingest_queue import IngestQueue\n"
                                   "assert IngestQueue(WorkerSettings.from_environment()).recover() >= 1\n"],
                                  cwd=Path(__file__).resolve().parents[2] / "worker", env=env,
                                  text=True, capture_output=True, timeout=90, check=False)
        assert recovery.returncode == 0, recovery.stderr[-1200:]
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
            assert db.execute("SELECT estado,error_codigo FROM prc.proceso_tarea WHERE proceso_id=%s",
                              (blocked_process_id,)).fetchone() == ("error", "preparacion_fallida")
            assert db.execute("SELECT estado FROM med.archivo WHERE id=%s", (third_id,)).fetchone()[0] == "error"

@pytest.mark.skipif(os.getenv('SYNTRA_RUN_INGEST_LOAD') != '1',
                    reason='Carga de rendimiento solo mediante activación explícita')
def test_ingest_performance_in_disposable_database(environment, monkeypatch):
    """Mide 100.000 o 1.000.000 de filas en PostgreSQL descartable."""
    from app.api.routers import measurement_files

    client, password, port = environment
    _close_previous_measurement_tasks(port)
    admin = _login(client, password)
    suffix = secrets.token_hex(6)
    with tempfile.TemporaryDirectory(prefix='syntra_ingest_load_') as storage:
        monkeypatch.setattr(measurement_files, 'settings',
                            replace(measurement_files.settings, upload_root=Path(storage)))
        with psycopg.connect(host='127.0.0.1', port=port, dbname='syntra_trial', user='postgres') as db:
            client_id = db.execute('INSERT INTO cfg.cliente(nombre) VALUES (%s) RETURNING id',
                                   (f'Cliente carga {suffix}',)).fetchone()[0]
            account_id = db.execute(
                "INSERT INTO cfg.cuenta(cliente_id,nombre,prompt_clusters,prompt_temas) "
                "VALUES (%s,%s,'clusters','temas') RETURNING id",
                (client_id, f'Cuenta carga {suffix}'),
            ).fetchone()[0]
        measurement = client.post('/api/v1/mediciones', headers=admin, json={
            'cuenta_id': account_id, 'nombre': f'Medición carga {suffix}',
            'fuentes': [{'origen': 'archivo', 'activo': True}],
        })
        assert measurement.status_code == 201, measurement.text
        measurement_id = measurement.json()['id']
        rows = int(os.getenv('SYNTRA_INGEST_LOAD_ROWS', '100000'))
        assert rows in (100_000, 1_000_000)
        file_ids = []
        for first in range(0, rows, 100_000):
            content = (b'url,published,content,source_type\n' + b''.join(
                 f'https://load.test/{suffix}/{index},07/10/26 12:46:00,Texto {index},'
                '"SOCIALMEDIA,SOCIALMEDIA_TWITTER"\n'.encode()
                for index in range(first, first + 100_000)
            ))
            assert len(content) <= 50 * 1024 * 1024
            reserved = client.post(f'/api/v1/mediciones/{measurement_id}/archivos/presignar',
                                   headers=admin, json={'nombre': f'load-{first}.csv',
                                                        'tamano_bytes': len(content),
                                                        'formato_id': 'talkwalker_simple'})
            assert reserved.status_code == 201, reserved.text
            file_id = reserved.json()['archivo_id']
            assert client.put(f'/api/v1/archivos/{file_id}/contenido', headers=admin,
                              content=content).status_code == 204
            file_ids.append(file_id)
        queued = client.post(f'/api/v1/mediciones/{measurement_id}/cargas/archivos',
                             headers={**admin, 'Idempotency-Key': f'load-{rows}-{suffix}-12345678'},
                             json={'archivo_ids': file_ids})
        assert queued.status_code == 202, queued.text
        process_id = queued.json()['proceso_id']
        worker_dir = Path(__file__).resolve().parents[2] / 'worker'
        common = {**os.environ, 'SYNTRA_WORKER_CONCURRENCY': '1',
                  'SYNTRA_DB_HOST': '127.0.0.1', 'SYNTRA_DB_PORT': str(port),
                  'SYNTRA_DB_NAME': 'syntra_trial'}
        prep = subprocess.run([sys.executable, '-c',
                               'from app.config import WorkerSettings\n'
                               'from app.infrastructure.measurement_queue import MeasurementQueue\n'
                               'q=MeasurementQueue(WorkerSettings.from_environment())\n'
                               't=q.claim("load:q.mantenimiento")\nassert t\n'
                               'q.process(t,"load:q.mantenimiento")\n'],
                              cwd=worker_dir, env={**common, 'SYNTRA_WORKER_QUEUE': 'q.mantenimiento',
                                                   'SYNTRA_WORKER_DB_USER': 'syntra_measurement_worker_login'},
                              text=True, capture_output=True, timeout=120, check=False)
        assert prep.returncode == 0, prep.stderr[-1000:]
        started = time.perf_counter()
        worker = subprocess.run([sys.executable, '-c',
                                 'from app.config import WorkerSettings\n'
                                 'from app.infrastructure.ingest_queue import IngestQueue\n'
                                 'q=IngestQueue(WorkerSettings.from_environment())\n'
                                 f'for _ in range({len(file_ids)}):\n'
                                 '    t=q.claim("load:q.ingesta")\n'
                                 '    assert t\n'
                                 '    q.process(t,"load:q.ingesta")\n'],
                                cwd=worker_dir,
                                env={**common, 'SYNTRA_WORKER_QUEUE': 'q.ingesta',
                                     'SYNTRA_WORKER_DB_USER': 'syntra_ingest_worker_login',
                                     'SYNTRA_UPLOAD_ROOT': storage},
                                text=True, capture_output=True, timeout=2100, check=False)
        wall_seconds = time.perf_counter() - started
        with psycopg.connect(host='127.0.0.1', port=port, dbname='syntra_trial', user='postgres') as db:
            state, duration_ms = db.execute('SELECT estado,duracion_ms FROM prc.proceso WHERE id=%s',
                                            (process_id,)).fetchone()
            persisted = db.execute('SELECT count(*) FROM dat.registro WHERE medicion_id=%s',
                                   (measurement_id,)).fetchone()[0]
        sla_seconds = 180 if rows == 100_000 else 1799
        print(f'INGEST_LOAD_{rows} wall_s={wall_seconds:.3f} process_ms={duration_ms} '
              f'rows={persisted} sla_s={sla_seconds} '
              f'met={duration_ms is not None and duration_ms <= sla_seconds * 1000}')
        assert worker.returncode == 0, worker.stderr[-1000:]
        assert state == 'completado' and persisted == rows


def test_measurement_tenant_isolation_for_api_processes_and_worker_task(environment):
    """Un usuario de una cuenta no puede crear, consultar ni ejecutar datos de otra."""
    client, password, port = environment
    _close_previous_measurement_tasks(port)
    admin_headers = _login(client, password)
    admin_id = client.get("/api/v1/auth/yo", headers=admin_headers).json()["usuario"]["id"]
    suffix = secrets.token_hex(6)
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
        db.execute("INSERT INTO cat.origen(id,codigo,nombre) VALUES (5,'manual','Manual de prueba') ON CONFLICT DO NOTHING")
        role_id = db.execute(
            "INSERT INTO seg.rol(nombre,descripcion,ver_todas,estado_id) VALUES (%s,'Rol acotado de prueba',false,1) RETURNING id",
            (f"Mediciones aisladas {suffix}",),
        ).fetchone()[0]
        db.execute("INSERT INTO seg.rol_permiso(rol_id,modulo,accion) VALUES "
                   "(%s,'mediciones','ver'),(%s,'mediciones','crear'),(%s,'procesos','ver')",
                   (role_id, role_id, role_id))
        scoped_user_id = db.execute(
            "INSERT INTO seg.usuario(nombre,email,password_hash,rol_id) VALUES (%s,%s,%s,%s) RETURNING id",
            ("Usuario aislado", f"aislado-{suffix}@syntra.local", password_hasher.hash(password), role_id),
        ).fetchone()[0]
        group_id = db.execute("INSERT INTO seg.grupo(nombre) VALUES (%s) RETURNING id",
                              (f"Grupo aislado {suffix}",)).fetchone()[0]
        client_a = db.execute("INSERT INTO cfg.cliente(nombre,creado_por,actualizado_por) VALUES (%s,%s,%s) RETURNING id",
                              (f"Cliente permitido {suffix}", admin_id, admin_id)).fetchone()[0]
        account_a = db.execute(
            "INSERT INTO cfg.cuenta(cliente_id,nombre,prompt_clusters,prompt_temas,creado_por,actualizado_por) "
            "VALUES (%s,%s,'clusters','temas',%s,%s) RETURNING id",
            (client_a, f"Cuenta permitida {suffix}", admin_id, admin_id),
        ).fetchone()[0]
        client_b = db.execute("INSERT INTO cfg.cliente(nombre,creado_por,actualizado_por) VALUES (%s,%s,%s) RETURNING id",
                              (f"Cliente ajeno {suffix}", admin_id, admin_id)).fetchone()[0]
        account_b = db.execute(
            "INSERT INTO cfg.cuenta(cliente_id,nombre,prompt_clusters,prompt_temas,creado_por,actualizado_por) "
            "VALUES (%s,%s,'clusters','temas',%s,%s) RETURNING id",
            (client_b, f"Cuenta ajena {suffix}", admin_id, admin_id),
        ).fetchone()[0]
        db.execute("INSERT INTO seg.grupo_usuario(grupo_id,usuario_id) VALUES (%s,%s)", (group_id, scoped_user_id))
        db.execute("INSERT INTO seg.grupo_cuenta(grupo_id,cuenta_id) VALUES (%s,%s)", (group_id, account_a))

    scoped_headers = _login(client, password, email=f"aislado-{suffix}@syntra.local")
    own = client.post("/api/v1/mediciones", headers=scoped_headers, json={
        "cuenta_id": account_a, "nombre": f"Medición permitida {suffix}",
        "fuentes": [{"origen": "manual", "activo": True}],
    })
    assert own.status_code == 201, own.text
    own_measurement = own.json()

    foreign = client.post("/api/v1/mediciones", headers=scoped_headers, json={
        "cuenta_id": account_b, "nombre": f"Intento ajeno {suffix}",
        "fuentes": [{"origen": "manual", "activo": True}],
    })
    assert foreign.status_code == 400
    assert "no está disponible" in foreign.json()["detail"]

    foreign_measurement = client.post("/api/v1/mediciones", headers=admin_headers, json={
        "cuenta_id": account_b, "nombre": f"Medición del otro cliente {suffix}",
        "fuentes": [{"origen": "manual", "activo": True}],
    })
    assert foreign_measurement.status_code == 201, foreign_measurement.text
    foreign_measurement_id = foreign_measurement.json()["id"]
    assert client.get(f"/api/v1/mediciones/{foreign_measurement_id}", headers=scoped_headers).status_code == 404
    listing = client.get("/api/v1/mediciones", headers=scoped_headers)
    assert listing.status_code == 200, listing.text
    assert [row["id"] for row in listing.json()["items"]] == [own_measurement["id"]]
    assert client.get(f"/api/v1/procesos/{foreign_measurement.json()['proceso_preparacion_id']}",
                      headers=scoped_headers).status_code == 404
    own_processes = client.get("/api/v1/procesos?tipo=preparar_medicion", headers=scoped_headers)
    assert own_processes.status_code == 200, own_processes.text
    assert {row["medicion_id"] for row in own_processes.json()} == {own_measurement["id"]}

    # Reclama la tarea ajena y prueba que la función de frontera rechaza un
    # cuenta_id alterado antes de crear particiones o completar el proceso.
    root = Path(__file__).resolve().parents[2]
    env = os.environ.copy()
    env.update({"SYNTRA_WORKER_QUEUE": "q.mantenimiento", "SYNTRA_WORKER_CONCURRENCY": "1",
                "SYNTRA_DB_HOST": "127.0.0.1", "SYNTRA_DB_PORT": str(port),
                "SYNTRA_DB_NAME": "syntra_trial", "SYNTRA_WORKER_DB_USER": "syntra_measurement_worker_login",
                "SYNTRA_WORKER_DB_PASSFILE": str(Path(tempfile.gettempdir()) / f"syntra-unset-passfile-{suffix}")})
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
        # Mantener en cola únicamente la tarea ajena para que el reclamo adversarial
        # no dependa del orden de creación de otras mediciones de esta misma prueba.
        own_process_id = own_measurement["proceso_preparacion_id"]
        db.execute("UPDATE prc.proceso_tarea SET estado='completado',finalizado_en=now(),ok=1 "
                   "WHERE proceso_id=%s", (own_process_id,))
        db.execute("UPDATE prc.proceso SET estado='completado',finalizado_en=now(),progreso=100 "
                   "WHERE id=%s", (own_process_id,))
    worker_id = f"trial:{suffix}:q.mantenimiento:isolation"
    adversarial = f"""import psycopg
from app.config import WorkerSettings
from app.infrastructure.measurement_queue import MeasurementQueue
settings = WorkerSettings.from_environment()
worker = {worker_id!r}
task = MeasurementQueue(settings).claim(worker)
assert task is not None and task.process_id == {foreign_measurement.json()['proceso_preparacion_id']}
with psycopg.connect(host='127.0.0.1', port={port}, dbname='syntra_trial', user='syntra_measurement_worker_login') as conn:
    try:
        conn.execute('SELECT ops.fn_preparar_medicion_segura(%s,%s,%s,%s,%s,%s)',
                     (task.id, worker, {account_a}, {foreign_measurement_id}, {admin_id}, 1))
    except psycopg.errors.InsufficientPrivilege as exc:
        assert exc.sqlstate == '42501'
        conn.rollback()
    else:
        raise AssertionError('La tarea aceptó el cuenta_id de otro cliente')
"""
    rejected = subprocess.run([sys.executable, "-c", adversarial], cwd=root / "worker", env=env,
                              capture_output=True, text=True, timeout=45, check=False)
    assert rejected.returncode == 0, rejected.stderr[-1600:]
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
        db.execute("UPDATE prc.proceso_tarea SET latido=now()-interval '10 minutes' WHERE proceso_id=%s",
                   (foreign_measurement.json()["proceso_preparacion_id"],))
    recover_code = f"""from app.config import WorkerSettings
from app.infrastructure.measurement_queue import MeasurementQueue
assert MeasurementQueue(WorkerSettings.from_environment()).recover() == 1
"""
    recovered = subprocess.run([sys.executable, "-c", recover_code], cwd=root / "worker", env=env,
                               capture_output=True, text=True, timeout=45, check=False)
    assert recovered.returncode == 0, recovered.stderr[-1600:]
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
        assert db.execute("SELECT estado FROM prc.proceso WHERE id=%s",
                          (foreign_measurement.json()["proceso_preparacion_id"],)).fetchone()[0] == "en_cola"
        assert db.execute(
            "SELECT count(*) FROM ops.fn_tablas_dat() t "
            "WHERE to_regclass('dat.'||quote_ident(t||'_c'||%s::text||'_m'||%s::text)) IS NOT NULL",
            (account_b, foreign_measurement_id),
        ).fetchone()[0] == 0


def test_measurement_migration_negative_preflight_rollback_and_retry(environment):
    """Preflight negativo y fallo tardío revierten todos los cambios transaccionales."""
    _client, _password, port = environment
    root = Path(__file__).resolve().parents[2]
    pg = Path("C:/Program Files/PostgreSQL/18/bin")
    migration = root / "database" / "entrega5_mediciones.sql"

    def snapshot():
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
            return db.execute("""SELECT
                (SELECT row_to_json(x)::text FROM (SELECT codigo,cola,sla_100k_s,sla_1m_s
                   FROM cat.tipo_proceso WHERE codigo='preparar_medicion') x),
                (SELECT row_to_json(x)::text FROM (SELECT operacion,max_intentos,backoff_base_ms,
                   backoff_max_ms,timeout_ms,usar_respaldo,notas FROM ops.politica_reintento
                   WHERE operacion='preparar_medicion') x),
                (SELECT proacl::text FROM pg_proc WHERE oid=
                   'ops.fn_preparar_medicion(integer,integer,integer[])'::regprocedure),
                has_function_privilege('syntra_measurement_worker_login',
                   'ops.fn_preparar_medicion(integer,integer,integer[])','EXECUTE'),
                has_function_privilege('syntra_measurement_worker_login',
                   'ops.fn_preparar_medicion_segura(bigint,text,integer,integer,integer,integer)','EXECUTE'),
                has_table_privilege('syntra_measurement_worker_login','ops.worker','DELETE')""").fetchone()

    before = snapshot()

    def run_migration(user: str, overrides: tuple[str, ...] = ()):
        return subprocess.run([
            str(pg / "psql.exe"), "-X", "-w", "-q", "-v", "ON_ERROR_STOP=1",
            *overrides, "-h", "127.0.0.1", "-p", str(port), "-U", user, "-d", "syntra_trial",
            "-f", str(migration),
        ], capture_output=True, text=True, timeout=60, check=False)

    negative_cases = (
        ("syntra_trial_deployer", ("-v", "SLA_1M_S=1799", "-v", "MED_TIMEOUT_MS=1799000"), "Falta -v SLA_100K_S"),
        ("syntra_trial_deployer", ("-v", "SLA_100K_S=0", "-v", "SLA_1M_S=1799", "-v", "MED_TIMEOUT_MS=1799000"), "enteros positivos"),
        ("postgres", ("-v", "SLA_100K_S=180", "-v", "SLA_1M_S=1799", "-v", "MED_TIMEOUT_MS=1799000"), "No ejecutar con superusuario"),
        ("syntra_api_login", ("-v", "SLA_100K_S=180", "-v", "SLA_1M_S=1799", "-v", "MED_TIMEOUT_MS=1799000"), "WITH GRANT OPTION"),
    )
    for user, overrides, expected in negative_cases:
        rejected = run_migration(user, overrides)
        output = rejected.stdout + rejected.stderr
        assert rejected.returncode != 0
        assert expected.casefold() in output.casefold(), output[-1600:]
        assert snapshot() == before

    # Inyectar una falla determinista después del preflight y de los GRANT/REVOKE
    # iniciales, dentro de la base descartable, para demostrar rollback transaccional.
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
        db.execute("""CREATE FUNCTION public.syntra_test_fail_measurement_migration()
                     RETURNS trigger LANGUAGE plpgsql AS $$BEGIN
                     RAISE EXCEPTION 'fallo inducido de prueba'; END$$""")
        db.execute("""CREATE TRIGGER syntra_test_fail_measurement_migration
                     BEFORE INSERT ON cat.tipo_proceso FOR EACH ROW
                     WHEN (NEW.codigo='preparar_medicion')
                     EXECUTE FUNCTION public.syntra_test_fail_measurement_migration()""")
    try:
        failed = run_migration("syntra_trial_deployer", ("-v", "SLA_100K_S=180", "-v", "SLA_1M_S=1799",
                                                         "-v", "MED_TIMEOUT_MS=1799000"))
        assert failed.returncode != 0
        assert "fallo inducido de prueba" in (failed.stdout + failed.stderr)
        assert snapshot() == before
    finally:
        with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
            db.execute("DROP TRIGGER IF EXISTS syntra_test_fail_measurement_migration ON cat.tipo_proceso")
            db.execute("DROP FUNCTION IF EXISTS public.syntra_test_fail_measurement_migration()")

    recovered = run_migration("syntra_trial_deployer", ("-v", "SLA_100K_S=180", "-v", "SLA_1M_S=1799",
                                                          "-v", "MED_TIMEOUT_MS=1799000"))
    assert recovered.returncode == 0, (recovered.stdout + recovered.stderr)[-2000:]
    with psycopg.connect(host="127.0.0.1", port=port, dbname="syntra_trial", user="postgres") as db:
        assert db.execute("SELECT sla_100k_s,sla_1m_s FROM cat.tipo_proceso WHERE codigo='preparar_medicion'").fetchone() == (180, 1799)
        assert db.execute("SELECT max_intentos,backoff_base_ms,backoff_max_ms,timeout_ms "
                          "FROM ops.politica_reintento WHERE operacion='preparar_medicion'").fetchone() == (3, 30_000, 120_000, 1_799_000)
