"""Serialización XLSX y almacenamiento privado local de reportes."""

from datetime import datetime, timedelta, timezone
import json
import os
from pathlib import Path
import re
import time
from uuid import uuid4
import zipfile

from openpyxl import Workbook

from app.private_files import restrict_to_current_user

MAX_SHEET_ROWS = 1_048_576
MAX_CELL = 30_000
FILE_NAME = re.compile(r"auditoria-[1-9][0-9]*(?:-[0-9a-f]{32})?\.(?:xlsx|tmp)\Z")
HEADERS = ("Fecha", "ID", "Usuario", "Módulo", "Acción", "Entidad", "ID de entidad", "Nombre de entidad",
           "Antes", "Después", "Detalle", "IP", "Solicitud")


def private_root(root: Path) -> Path:
    if not root.is_absolute():
        raise ValueError("SYNTRA_REPORT_ROOT debe ser una ruta absoluta privada")
    if root.resolve(strict=False).is_relative_to(Path(__file__).resolve().parents[3]):
        raise ValueError("SYNTRA_REPORT_ROOT debe estar fuera del repositorio")
    if root.is_symlink():
        raise ValueError("SYNTRA_REPORT_ROOT no puede ser un enlace")
    root.mkdir(mode=0o700, parents=True, exist_ok=True)
    restrict_to_current_user(root)
    return root.resolve(strict=True)


def safe_cell(value: object) -> str:
    if value is None:
        return ""
    raw = str(value)
    return "'" + raw if raw.lstrip().startswith(("=", "+", "-", "@")) else raw


class ReportFiles:
    def __init__(self, root: Path):
        self.root = root

    def path(self, process_id: int, extension: str) -> Path:
        if process_id < 1 or extension not in {"xlsx", "tmp"}:
            raise ValueError("Archivo de proceso inválido")
        return self.root / f"auditoria-{process_id}.{extension}"

    def unique_temp(self, process_id: int) -> Path:
        return self.root / f"auditoria-{process_id}-{uuid4().hex}.tmp"

    def valid_final(self, process_id: int) -> bool:
        final = self.path(process_id, "xlsx")
        return final.is_file() and final.stat().st_size > 0 and zipfile.is_zipfile(final)

    def write(self, process_id: int, rows, staged_path: Path | None = None,
              deadline: float | None = None) -> int | tuple[int, Path | None]:
        temp = staged_path or self.path(process_id, "tmp")
        final = self.path(process_id, "xlsx")
        if self.valid_final(process_id):
            count = 0
            for _ in rows:
                if deadline is not None and time.monotonic() >= deadline:
                    raise TimeoutError("Se agotó el tiempo de generación")
                count += 1
            return (count, None) if staged_path is not None else count
        workbook = Workbook(write_only=True)
        sheet_index = 1
        sheet = workbook.create_sheet(f"Eventos {sheet_index}")
        sheet.append(HEADERS)
        sheet_rows = 1
        content_sheet = None
        content_rows = 0
        content_index = 0
        count = 0

        def content(date: str, identifier: int, field: str, raw: str) -> str:
            nonlocal content_sheet, content_rows, content_index
            for part, start in enumerate(range(0, len(raw), MAX_CELL), start=1):
                if content_sheet is None or content_rows >= MAX_SHEET_ROWS:
                    content_index += 1
                    content_sheet = workbook.create_sheet(f"Contenido {content_index}")
                    content_sheet.append(("Fecha", "ID", "Campo", "Parte", "Texto"))
                    content_rows = 1
                content_sheet.append((date, identifier, field, part, safe_cell(raw[start:start + MAX_CELL])))
                content_rows += 1
            return f"Ver Contenido ({field}, {len(raw)} caracteres)"

        for row in rows:
            if deadline is not None and time.monotonic() >= deadline:
                raise TimeoutError("Se agotó el tiempo de generación")
            if sheet_rows >= MAX_SHEET_ROWS:
                sheet_index += 1
                sheet = workbook.create_sheet(f"Eventos {sheet_index}")
                sheet.append(HEADERS)
                sheet_rows = 1
            before = json.dumps(row["antes"], ensure_ascii=False, default=str) if row["antes"] is not None else ""
            after = json.dumps(row["despues"], ensure_ascii=False, default=str) if row["despues"] is not None else ""
            date = row["fecha"].isoformat()
            values = [date, row["id"], row["usuario_id"], row["modulo"], row["accion"],
                      row["entidad_tipo"], row["entidad_id"], row["entidad_nombre"],
                      before, after, row["detalle"], row["ip"], row["request_id"]]
            for index, value in enumerate(values):
                if value is not None and len(str(value)) > MAX_CELL:
                    values[index] = content(date, row["id"], HEADERS[index], str(value))
            sheet.append(tuple(safe_cell(value) for value in values))
            sheet_rows += 1
            count += 1
        if deadline is not None and time.monotonic() >= deadline:
            raise TimeoutError("Se agotó el tiempo de generación")
        workbook.save(temp)
        if deadline is not None and time.monotonic() >= deadline:
            raise TimeoutError("Se agotó el tiempo de generación")
        if staged_path is not None:
            return count, temp
        os.replace(temp, final)
        return count

    def publish(self, process_id: int, staged: Path | None) -> Path:
        final = self.path(process_id, "xlsx")
        if staged is not None:
            if staged.parent != self.root or not re.fullmatch(
                rf"auditoria-{process_id}-[0-9a-f]{{32}}\.tmp", staged.name
            ) or staged.is_symlink():
                raise ValueError("Archivo temporal de exportación inválido")
            if self.valid_final(process_id):
                staged.unlink(missing_ok=True)
            else:
                os.replace(staged, final)
        if not self.valid_final(process_id):
            raise FileNotFoundError("El reporte final no está disponible")
        return final

    def cleanup(self) -> tuple[int, int]:
        now = datetime.now(timezone.utc)
        removed_final = removed_temp = 0
        for path in self.root.iterdir():
            if not path.is_file() or path.is_symlink() or not FILE_NAME.fullmatch(path.name):
                continue
            age = now - datetime.fromtimestamp(path.stat().st_mtime, timezone.utc)
            if path.suffix == ".xlsx" and age > timedelta(days=7):
                path.unlink()
                removed_final += 1
            elif path.suffix == ".tmp" and age > timedelta(hours=24):
                path.unlink()
                removed_temp += 1
        return removed_final, removed_temp
