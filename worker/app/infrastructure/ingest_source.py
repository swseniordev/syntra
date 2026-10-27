"""Adaptador local de archivos privados para q.ingesta."""

from __future__ import annotations

import csv
import hashlib
import hmac
import os
from pathlib import Path
import re
from typing import Any, Iterator

from openpyxl import load_workbook


KEY_PATTERN = re.compile(r"[1-9][0-9]*/[1-9][0-9]*/[0-9a-f]{32}(?:-[0-9a-f]{64})?\.(?:csv|xlsx|xls)\Z")
REQUIRED_HEADERS = frozenset({"url", "published"})


def _validate_headers(headers: list[Any]) -> None:
    if not headers or any(not isinstance(value, str) or not value for value in headers) or len(set(headers)) != len(headers):
        raise ValueError("El archivo no tiene cabeceras únicas y válidas")
    missing = REQUIRED_HEADERS.difference(headers)
    if missing:
        raise ValueError("Faltan columnas obligatorias: " + ", ".join(sorted(missing)))


def rows_from_file(path: Path, extension: str, encoding: str, delimiter: str | None) -> Iterator[dict[str, Any]]:
    if extension == "csv":
        with path.open("r", encoding=encoding, newline="") as stream:
            reader = csv.DictReader(stream, delimiter=delimiter or ",")
            _validate_headers(reader.fieldnames or [])
            yield from reader
    elif extension == "xlsx":
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            sheet = book.active
            iterator = sheet.iter_rows(values_only=True)
            headers = next(iterator, None)
            _validate_headers(list(headers or []))
            for cells in iterator:
                yield dict(zip(headers, cells))
        finally:
            book.close()
    else:
        raise ValueError("Extensión no admitida para ingesta")


class LocalIngestSource:
    def __init__(self, root: Path):
        self.root = root.resolve(strict=False)

    def rows(self, key: str, extension: str, encoding: str,
             delimiter: str | None) -> Iterator[dict[str, Any]]:
        path = self._path(key)
        yield from rows_from_file(path, extension, encoding, delimiter)

    def verify(self, key: str, expected_sha256: bytes) -> None:
        digest = hashlib.sha256()
        with self._path(key).open('rb') as stream:
            for chunk in iter(lambda: stream.read(1024 * 1024), b''):
                digest.update(chunk)
        if not hmac.compare_digest(digest.digest(), expected_sha256):
            raise ValueError('El contenido del archivo ya no coincide con la carga validada')

    def _path(self, key: str) -> Path:
        if not KEY_PATTERN.fullmatch(key):
            raise ValueError("Archivo privado inválido")
        path = self.root.joinpath(*key.split("/"))
        if os.name == "nt":
            if not path.is_relative_to(self.root):
                raise ValueError("Archivo privado inválido")
            current = path
            while current != self.root:
                if current.is_symlink() or current.is_junction():
                    raise ValueError("Archivo privado inválido")
                current = current.parent
        else:
            path = path.resolve(strict=True)
            if not path.is_relative_to(self.root) or path.is_symlink():
                raise ValueError("Archivo privado inválido")
        if not path.is_file():
            raise FileNotFoundError("Archivo privado no disponible")
        return path
